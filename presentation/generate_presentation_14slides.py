"""VisionInsight 기술 발표 자료 14장 생성 스크립트
(14-Slide Presentation with Core Engineering Challenges)

목적:
- 기존 13장 구성(Slide 1~8 및 기존 Slide 9~13)을 100% 완전 보존
- Slide 8(구조화 출력 개선)과 기존 Slide 9(실서버 E2E 결과) 사이에
  신규 Slide 9("핵심 기술 난제와 해결 과정") 삽입 ➔ 총 14장 구성
- 신규 Slide 9 필수 포함 3대 기술 난제:
  1. Gemini 구조화 출력 HTTP 400 (복잡한 스키마 거부 ➔ 단순 JSON + Python 검증)
  2. Gemini 최종 JSON HTTP 504 (최종 생성 단계 타임아웃 ➔ 요청별 timeout 분리 실측)
  3. ERROR 보고서의 근거 검증 공백 (오류 시나리오 누락 ➔ 17개 필수 사실 정책 보완, Mock PASS, 실서버 NOT TESTED)
- 최종 산출물: presentation/VisionInsight_Presentation_14slides.pptx
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation_14slides():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette (Industrial Dark Navy & High Contrast Accents)
    BG_COLOR = RGBColor(12, 19, 34)          # #0C1322
    CARD_BG = RGBColor(22, 30, 46)           # #161E2E
    CARD_BG_ALT = RGBColor(16, 23, 38)       # #101726
    CARD_BORDER = RGBColor(42, 54, 73)       # #2A3649
    PRIMARY_CYAN = RGBColor(76, 215, 246)    # #4CD7F6
    SUCCESS_GREEN = RGBColor(52, 211, 153)   # #34D399
    WARNING_AMBER = RGBColor(251, 191, 36)   # #FBBF24
    DANGER_CORAL = RGBColor(248, 113, 113)   # #F87171
    TEXT_WHITE = RGBColor(248, 250, 252)     # #F8FAFC
    TEXT_LIGHT = RGBColor(203, 213, 225)     # #CBD5E1
    TEXT_MUTED = RGBColor(148, 163, 184)     # #94A3B8

    FONT_FAMILY = "Malgun Gothic"

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR
        return bg

    def add_header(slide, title_text, category_text="VISIONINSIGHT  |  INDUSTRIAL VISION AI AGENT", subtitle_text=None):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.36), Inches(11.7), Inches(0.25))
        tf = cat_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = category_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_CYAN

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.62), Inches(11.7), Inches(0.68))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_top = tf_t.margin_bottom = tf_t.margin_left = tf_t.margin_right = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(20)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

        if subtitle_text:
            p_sub = tf_t.add_paragraph()
            p_sub.text = subtitle_text
            p_sub.font.name = FONT_FAMILY
            p_sub.font.size = Pt(11)
            p_sub.font.color.rgb = TEXT_LIGHT
            p_sub.space_before = Pt(2)

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    add_card(s1, Inches(0.8), Inches(1.1), Inches(11.733), Inches(5.4), bg_color=CARD_BG_ALT, border_color=PRIMARY_CYAN)

    tbox = s1.shapes.add_textbox(Inches(1.4), Inches(1.7), Inches(10.5), Inches(2.2))
    tf = tbox.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "VisionInsight"
    p0.font.name = FONT_FAMILY
    p0.font.size = Pt(46)
    p0.font.bold = True
    p0.font.color.rgb = PRIMARY_CYAN

    p1 = tf.add_paragraph()
    p1.text = "AI 기반 병 외관 이상 탐지 및 근거 기반 진단 시스템"
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_before = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "MVTec Bottle 기반 PatchCore 이상 탐지와 Gemini Function Calling의 결합"
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(13)
    p2.font.color.rgb = TEXT_LIGHT
    p2.space_before = Pt(10)

    bbox = s1.shapes.add_textbox(Inches(1.4), Inches(4.5), Inches(10.5), Inches(1.6))
    tf_b = bbox.text_frame
    tf_b.word_wrap = True
    bp0 = tf_b.paragraphs[0]
    bp0.text = "핵심 기술 스택: Python 3.11  |  FastAPI  |  PyTorch/PatchCore  |  Gemini 3.8 Flash (Function Calling)"
    bp0.font.name = FONT_FAMILY
    bp0.font.size = Pt(12)
    bp0.font.bold = True
    bp0.font.color.rgb = PRIMARY_CYAN

    bp1 = tf_b.add_paragraph()
    bp1.text = "검증 완료: 단일 이미지 CPU 추론 평균 42.53ms  |  E2E 실서버 검증 (T01 양품 PASS / T03 경고 PASS)  |  전체 pytest 886 PASS"
    bp1.font.name = FONT_FAMILY
    bp1.font.size = Pt(11)
    bp1.font.color.rgb = TEXT_MUTED
    bp1.space_before = Pt(6)

    bp2 = tf_b.add_paragraph()
    bp2.text = "발표일: 2026년 10월 20일  |  발표자: VisionInsight 개발팀"
    bp2.font.name = FONT_FAMILY
    bp2.font.size = Pt(11)
    bp2.font.color.rgb = TEXT_LIGHT
    bp2.space_before = Pt(10)

    # ==========================================
    # SLIDE 2: Problem Definition
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "산업 현장의 과제: 규칙 기반 검사의 한계와 AI 진단의 필요성")

    add_card(s2, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.3), border_color=DANGER_CORAL)
    tb_left2 = s2.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_l2 = tb_left2.text_frame
    tf_l2.word_wrap = True
    p = tf_l2.paragraphs[0]
    p.text = "기존 룰베이스 머신비전의 한계"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = DANGER_CORAL

    problems_final = [
        ("미세 결함 검출 한계", "고정 임계값 기반 검사로는 비정형 파손, 오염 등의 미세 이상 탐지 어려움"),
        ("환경 변화에 따른 오판정", "조명 밝기 저하나 카메라 노출 변화 시 정상 제품을 불량(NG)으로 오판정"),
        ("이유 없는 판정 결과", "단순히 Pass/Fail 수치만 출력되어, 작업자가 재검사 필요 여부나 점검 항목을 파악하기 어려움")
    ]
    for title, desc in problems_final:
        p = tf_l2.add_paragraph()
        p.text = "• " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(12)
        p_desc = tf_l2.add_paragraph()
        p_desc.text = "   " + desc
        p_desc.font.name = FONT_FAMILY
        p_desc.font.size = Pt(10.5)
        p_desc.font.color.rgb = TEXT_LIGHT
        p_desc.space_before = Pt(2)

    add_card(s2, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.3), border_color=PRIMARY_CYAN)
    tb_right2 = s2.shapes.add_textbox(Inches(7.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_r2 = tb_right2.text_frame
    tf_r2.word_wrap = True
    p = tf_r2.paragraphs[0]
    p.text = "VisionInsight가 제시하는 해결책"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_CYAN

    solutions_final = [
        ("PatchCore 비지도 이상 탐지", "정상 이미지 180장을 기반으로 메모리 뱅크를 구축하고, 독립 정상 29장 통계로 임계값(34.1442)을 산출하여 미세 이상 점수화"),
        ("품질 경고와 결함 판정 분리", "조도 저하 등 영상 품질 경고 시 제품 결함(NG)과 구분하여 서버 최종 판정을 '재검사 필요(REINSPECT)'로 분류"),
        ("Gemini 도구 기반 원인 진단", "Function Calling으로 조도/로그를 사후 조회하여 추정 원인, 현장 점검 권고와 안전 한계를 구조화 보고")
    ]
    for title, desc in solutions_final:
        p = tf_r2.add_paragraph()
        p.text = "✔ " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = SUCCESS_GREEN
        p.space_before = Pt(12)
        p_desc = tf_r2.add_paragraph()
        p_desc.text = "   " + desc
        p_desc.font.name = FONT_FAMILY
        p_desc.font.size = Pt(10.5)
        p_desc.font.color.rgb = TEXT_LIGHT
        p_desc.space_before = Pt(2)

    # ==========================================
    # SLIDE 3: System Architecture
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "시스템 아키텍처: Vision, Decision, Agent의 명확한 역할 분리")

    layers_final = [
        ("Layer 1: Vision 계층", PRIMARY_CYAN, [
            ("입력 영상 전처리", "900×900 이미지 수신 및 영상 통계(평균 밝기, 표준편차) 계산"),
            ("품질 상태 평가", "밝기 정상(NORMAL), 경고(WARNING_LOW_BRIGHTNESS), 실패(ERROR) 분류"),
            ("PatchCore 추론", "ResNet-18 특징 추출 및 이상 점수 산출 (개발 환경 CPU 평균 42.53ms)")
        ]),
        ("Layer 2: Decision 계층 (서버 규칙)", WARNING_AMBER, [
            ("규칙 기반 종합 판정", "영상 품질과 모델 결과를 종합하여 final_result 결정 (우선순위 1~6단계 규칙)"),
            ("REINSPECT 분리 판정", "영상 품질 경고 시 서버의 최종 판정을 REINSPECT로 분류하여 결함 판정과 구분"),
            ("판정 불변 원칙", "서버가 확정한 판정값은 이후 Agent가 절대 변경할 수 없음 (Gemini 판정 변경 차단)")
        ]),
        ("Layer 3: Agent 계층 (Gemini AI)", SUCCESS_GREEN, [
            ("도구 기반 정보 조회", "Function Calling으로 검사 결과 및 품질 메트릭을 사후 조회"),
            ("근거 기반 진단 보고서", "추정 원인 가설, 현장 점검 권고, 안전 한계를 구조화하여 제안"),
            ("Evidence 검증", "모델의 모든 관측 사실(Fact)을 서버 메모리에서 1:1 결정적 검증")
        ])
    ]
    card_width = Inches(3.7)
    cx = 0.8
    for ltitle, lcol, litems in layers_final:
        add_card(s3, Inches(cx), Inches(1.4), card_width, Inches(4.35), border_color=lcol)
        tb_l = s3.shapes.add_textbox(Inches(cx + 0.2), Inches(1.55), card_width - Inches(0.4), Inches(4.0))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True
        p = tf_l.paragraphs[0]
        p.text = ltitle
        p.font.name = FONT_FAMILY
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = lcol

        for ititle, idesc in litems:
            p_t = tf_l.add_paragraph()
            p_t.text = "• " + ititle
            p_t.font.name = FONT_FAMILY
            p_t.font.size = Pt(11)
            p_t.font.bold = True
            p_t.font.color.rgb = TEXT_WHITE
            p_t.space_before = Pt(8)

            p_d = tf_l.add_paragraph()
            p_d.text = "   " + idesc
            p_d.font.name = FONT_FAMILY
            p_d.font.size = Pt(9.5)
            p_d.font.color.rgb = TEXT_LIGHT
            p_d.space_before = Pt(2)
        cx += 4.0

    add_card(s3, Inches(0.8), Inches(5.95), Inches(11.733), Inches(0.8), bg_color=CARD_BG_ALT, border_color=PRIMARY_CYAN)
    tb_c3 = s3.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(11.3), Inches(0.7))
    tf_c3 = tb_c3.text_frame
    tf_c3.word_wrap = True
    p_c3 = tf_c3.paragraphs[0]
    p_c3.text = "💡 아키텍처 원칙: Vision(고속 추론) → Decision(서버 판정 확정) → Agent(사후 원인 진단)으로 책임을 분리하며, 현재는 서버 판정값을 확정하는 프로토타입으로 실제 물리적 PLC·컨베이어 제어는 수행하지 않습니다."
    p_c3.font.name = FONT_FAMILY
    p_c3.font.size = Pt(10.5)
    p_c3.font.color.rgb = TEXT_WHITE
    p_c3.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 4: PatchCore Vision Performance
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Vision 계층 성능: PatchCore 이상 탐지 실측 결과")

    add_card(s4, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.3))
    tb_m = s4.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    p = tf_m.paragraphs[0]
    p.text = "핵심 성능 지표 (83장 테스트셋)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_CYAN

    metrics_final = [
        ("AUROC", "0.8286", "이상 탐지 변별력 (전체 면적 기준)"),
        ("정밀도 (Precision)", "100.0%", "이상 감지 시 결함 신뢰도 (TP: 32건, FP: 0건 기준)"),
        ("재현율 (Recall)", "50.79%", "미탐 존재: 63건 중 32건 탐지 (TP: 32건, FN: 31건 기준)"),
        ("고정 임계값", "34.1442", "독립 정상 검증 데이터 29장의 통계 기반 (μ + 3σ = 34.1442)"),
        ("CPU 단일 추론 시간", "평균 42.53ms", "개발 환경(x86_64 CPU)의 단일 이미지 평균 추론 시간 (T01: 41.3ms, T03: 42.7ms)")
    ]
    for name, val, desc in metrics_final:
        p = tf_m.add_paragraph()
        p.text = f"{name}: {val}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = SUCCESS_GREEN if ("100" in val or "0.82" in val) else WARNING_AMBER if ("50" in val) else TEXT_WHITE
        p.space_before = Pt(8)
        p_d = tf_m.add_paragraph()
        p_d.text = f"  ({desc})"
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_MUTED

    add_card(s4, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.3))
    tb_cm = s4.shapes.add_textbox(Inches(7.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_cm = tb_cm.text_frame
    tf_cm.word_wrap = True
    p = tf_cm.paragraphs[0]
    p.text = "혼동 행렬 및 엔지니어링 분석"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_CYAN

    cm_rows_final = [
        ("진양성 (TP)", "32건", "실제 결함을 이상(NG)으로 정확히 탐지"),
        ("위양성 (FP)", "0건", "정상 제품을 결함으로 오판정하지 않음 (정밀도 100%)"),
        ("진음성 (TN)", "20건", "정상 제품 20장을 양품으로 정확히 판정"),
        ("위음성 (FN)", "31건", "미세 파손/오염 등 31건 미탐지 (재현율 50.79%)")
    ]
    for title, cnt, desc in cm_rows_final:
        p = tf_cm.add_paragraph()
        p.text = f"• {title}: {cnt}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = DANGER_CORAL if "FN" in title else SUCCESS_GREEN
        p.space_before = Pt(8)
        p_d = tf_cm.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_LIGHT

    p_note4 = tf_cm.add_paragraph()
    p_note4.text = "💡 시사점: FP 0건으로 과검은 없으나, 미탐(FN 31건) 한계가 존재하므로 모델의 OK 판정이 실제 무결점을 보증하지 않는다는 점을 안전 가이드에 필수 명시함. (정밀도 100%를 산업 현장 검출 성능 보장으로 확대 해석 금지)"
    p_note4.font.name = FONT_FAMILY
    p_note4.font.size = Pt(9.5)
    p_note4.font.color.rgb = WARNING_AMBER
    p_note4.space_before = Pt(12)

    # ==========================================
    # SLIDE 5: Decision Logic & REINSPECT
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Decision 계층: 대표 판정 규칙 4가지와 REINSPECT 분류")

    rules_final = [
        ("Rule 1: 파이프라인 실패 ➔ ERROR", DANGER_CORAL, "status == 'FAILED'", "INSPECTION_FAILED", "영상 로드/포맷 오류 등 검사 파이프라인 실패 시 ERROR 처리 및 오류 이벤트 기록"),
        ("Rule 2: 영상 품질 경고 ➔ REINSPECT", WARNING_AMBER, "image_quality_status != 'NORMAL'", "IMAGE_QUALITY_WARNING", "조도 저하(T03) 시 모델이 OK여도 제품 결함(NG)과 구분하여 '재검사 필요'로 분류"),
        ("Rule 4: 제품 이상 탐지 ➔ NG", DANGER_CORAL, "품질 NORMAL and result == 'NG'", "PRODUCT_ANOMALY_DETECTED", "PatchCore 이상 점수가 임계값을 초과하여 모델 NG로 판정되고, 서버 규칙에 따라 최종 NG로 분류"),
        ("Rule 5: 정상 검사 통과 ➔ OK", SUCCESS_GREEN, "품질 NORMAL and result == 'OK'", "ALL_CHECKS_PASSED", "영상 품질 NORMAL이고 이상 점수가 임계값(34.1442) 이하인 정상 제품 (T01)")
    ]
    ry = 1.4
    for rtitle, rcol, rsub, rcode, rdesc in rules_final:
        add_card(s5, Inches(0.8), Inches(ry), Inches(11.733), Inches(1.05), border_color=rcol)
        tb_r = s5.shapes.add_textbox(Inches(1.1), Inches(ry + 0.08), Inches(11.1), Inches(0.9))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True
        p = tf_r.paragraphs[0]
        p.text = f"{rtitle}  [{rcode}]"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = rcol

        p_desc = tf_r.add_paragraph()
        p_desc.text = f"• 조건: {rsub}  ➔  설명: {rdesc}"
        p_desc.font.name = FONT_FAMILY
        p_desc.font.size = Pt(10.2)
        p_desc.font.color.rgb = TEXT_LIGHT
        p_desc.space_before = Pt(2)

        ry += 1.15

    add_card(s5, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.8), bg_color=CARD_BG_ALT, border_color=PRIMARY_CYAN)
    tb_c5 = s5.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.7))
    tf_c5 = tb_c5.text_frame
    tf_c5.word_wrap = True
    p_c5 = tf_c5.paragraphs[0]
    p_c5.text = "💡 대표 판정 규칙 4가지 (전체 6개 규칙 중 우선순위 반영): Rule 3(모델 미수행/UNKNOWN ➔ DEFECT_DETECTION_UNAVAILABLE), Rule 6(미정의 조합 ➔ UNKNOWN_STATE_COMBINATION) 포함. 모델 소견(result)과 서버 최종 판정(final_result)을 엄격히 분리하며, REINSPECT는 물리적 폐기/격리 제어가 아닌 서버 판정 상태 분류입니다."
    p_c5.font.name = FONT_FAMILY
    p_c5.font.size = Pt(10.0)
    p_c5.font.color.rgb = TEXT_WHITE
    p_c5.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 6: Gemini Function Calling
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "AI Agent: Gemini Function Calling 기반 정밀 조사 체계")

    tools_final = [
        ("1. get_inspection_result", "검사 기본 상태 조회", "검사 상태, 모델 판정, 최종 판정, 사유 코드 조회 (E0001)"),
        ("2. analyze_image_quality", "영상 품질 세부 분석", "평균 밝기, 표준편차, min/max 동적범위 조회 (E0002)"),
        ("3. get_execution_logs", "실행 감사 로그 확인", "시각순 이벤트 로그(IMAGE_LOAD_FAILED 등) 추적"),
        ("4. get_inspection_history", "과거 이력 조회", "동일 제품의 직전 검사 이력 트렌드 확인")
    ]
    col_w = Inches(5.7)
    for idx, (tname, tsub, tdesc) in enumerate(tools_final):
        col_idx = idx % 2
        row_idx = idx // 2
        lx = 0.8 if col_idx == 0 else 6.8
        ly = 1.4 + row_idx * 1.85
        add_card(s6, Inches(lx), Inches(ly), col_w, Inches(1.65))
        tb_t = s6.shapes.add_textbox(Inches(lx + 0.3), Inches(ly + 0.15), col_w - Inches(0.6), Inches(1.35))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True
        p = tf_t.paragraphs[0]
        p.text = tname
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_CYAN

        p_sub = tf_t.add_paragraph()
        p_sub.text = "• 목적: " + tsub
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(10.5)
        p_sub.font.bold = True
        p_sub.font.color.rgb = TEXT_WHITE
        p_sub.space_before = Pt(3)

        p_d = tf_t.add_paragraph()
        p_d.text = "• 내용: " + tdesc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_before = Pt(2)

    add_card(s6, Inches(0.8), Inches(5.3), Inches(11.733), Inches(1.4), bg_color=CARD_BG_ALT, border_color=SUCCESS_GREEN)
    tb_flow = s6.shapes.add_textbox(Inches(1.0), Inches(5.4), Inches(11.3), Inches(1.2))
    tf_f = tb_flow.text_frame
    tf_f.word_wrap = True
    p = tf_f.paragraphs[0]
    p.text = "도구 실행 파이프라인 (보안 및 책임 경계)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    p_fdesc = tf_f.add_paragraph()
    p_fdesc.text = "Gemini는 Function Calling으로 조회 도구와 인자를 제안하고, FastAPI AgentService가 호출을 검증한 뒤 허용된 Python 조회 도구를 서버에서 실행합니다. (제어 흐름: 1.Gemini 도구 제안 ➔ 2.서버 인자 검증 ➔ 3.Python 실행 ➔ 4.FunctionResponse 전달 ➔ 5.최종 단순 JSON 보고서 생성 ➔ 6.Pydantic 및 Evidence 검증)"
    p_fdesc.font.name = FONT_FAMILY
    p_fdesc.font.size = Pt(9.8)
    p_fdesc.font.color.rgb = TEXT_LIGHT
    p_fdesc.space_before = Pt(3)

    # ==========================================
    # SLIDE 7: Evidence Trace
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Evidence 메커니즘: 결정적 근거 검증을 통한 신뢰성 확보")

    ev_pillars_final = [
        ("Pillar 1: Evidence ID", PRIMARY_CYAN, "서버 발급 도구 영수증", "도구 실행 시 서버가 E0001, E0002 고유 ID를 부여하여 모델의 허위 실행 주장 차단"),
        ("Pillar 2: Fact ID & Pointer", SUCCESS_GREEN, "JSON Pointer 1:1 결정적 검증", "보고서에 포함된 구조화된 관측 사실(F0001~)을 서버의 실제 도구 결과(/data/final_result 등)와 결정적으로 대조"),
        ("Pillar 3: 근거 앵커링", WARNING_AMBER, "권장 조치-근거 필수 연결", "모든 추천 조치(recommended_checks)는 반드시 유효한 관측 사실(evidence_refs)을 참조해야 통과"),
        ("Pillar 4: 판정 불변성", DANGER_CORAL, "서버 확정 판정 보호", "Gemini는 사후 진단 보조 역할만 수행하며, 서버가 확정한 final_result를 절대 변경할 수 없음")
    ]
    ey = 1.4
    for ptitle, pcol, psub, pdesc in ev_pillars_final:
        add_card(s7, Inches(0.8), Inches(ey), Inches(11.733), Inches(1.05), border_color=pcol)
        tb_p = s7.shapes.add_textbox(Inches(1.1), Inches(ey + 0.08), Inches(11.1), Inches(0.9))
        tf_p = tb_p.text_frame
        tf_p.word_wrap = True
        p = tf_p.paragraphs[0]
        p.text = f"{ptitle}  ({psub})"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = pcol

        p_d = tf_p.add_paragraph()
        p_d.text = "• 동작 원리: " + pdesc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10.2)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_before = Pt(2)

        ey += 1.15

    add_card(s7, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.8), bg_color=CARD_BG_ALT, border_color=SUCCESS_GREEN)
    tb_c7 = s7.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.7))
    tf_c7 = tb_c7.text_frame
    tf_c7.word_wrap = True
    p_c7 = tf_c7.paragraphs[0]
    p_c7.text = "💡 실제 실서버 검증 성과: T01의 관측 사실 6개와 T03의 관측 사실 12개가 서버의 Evidence 검증을 통과했습니다. (서버에 등록된 필드에 대해 엄밀한 1:1 결정적 일치를 확인하며, 모든 자연어 가설이나 미래의 출력을 완전 검증한다고 과장하지 않음)"
    p_c7.font.name = FONT_FAMILY
    p_c7.font.size = Pt(10.0)
    p_c7.font.color.rgb = TEXT_WHITE
    p_c7.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 8: Structured Output Compatibility
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "기술적 문제 해결: Gemini 구조화 출력 호환성 개선")

    add_card(s8, Inches(0.8), Inches(1.4), Inches(5.7), Inches(4.3), border_color=DANGER_CORAL)
    tb_b8 = s8.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(5.1), Inches(3.8))
    tf_b8 = tb_b8.text_frame
    tf_b8.word_wrap = True
    p = tf_b8.paragraphs[0]
    p.text = "Before: 직접 구조화 출력 시도 ➔ HTTP 400"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = DANGER_CORAL

    b_points_final = [
        ("복잡한 Pydantic 스키마 전달", "DiagnosisDraft 스키마를 response_json_schema로 직접 전송"),
        ("실서버 HTTP 400 거부", "Gemini 서버가 스키마 호환성 문제로 ClientError 반환 (정확한 거부 속성은 UNKNOWN)"),
        ("결과", "사후 진단 보고서 생성 실패")
    ]
    for title, desc in b_points_final:
        p = tf_b8.add_paragraph()
        p.text = "• " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)
        p_d = tf_b8.add_paragraph()
        p_d.text = "   " + desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT

    add_card(s8, Inches(6.8), Inches(1.4), Inches(5.7), Inches(4.3), border_color=SUCCESS_GREEN)
    tb_a8 = s8.shapes.add_textbox(Inches(7.1), Inches(1.65), Inches(5.1), Inches(3.8))
    tf_a8 = tb_a8.text_frame
    tf_a8.word_wrap = True
    p = tf_a8.paragraphs[0]
    p.text = "After: 스키마 단순화 + Python 검증 ➔ PASS"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    a_points_final = [
        ("출력 스키마 단순화", "Gemini에는 단순화된 JSON 스키마만 전송 (토큰 선택 지시)"),
        ("토큰 카탈로그 주입", "서버가 수집한 사실(K0001~)을 프롬프트로 전달, 모델은 토큰만 선택"),
        ("Python 엄격 검증 파이프라인", "Gemini 출력 ➔ Python DiagnosisDraft 변환 ➔ Pydantic ➔ 결정적 Evidence 검증")
    ]
    for title, desc in a_points_final:
        p = tf_a8.add_paragraph()
        p.text = "✔ " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_CYAN
        p.space_before = Pt(10)
        p_d = tf_a8.add_paragraph()
        p_d.text = "   " + desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT

    add_card(s8, Inches(0.8), Inches(5.95), Inches(11.733), Inches(0.8), bg_color=CARD_BG_ALT, border_color=PRIMARY_CYAN)
    tb_sum8 = s8.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(11.3), Inches(0.7))
    tf_s8 = tb_sum8.text_frame
    p_s8 = tf_s8.paragraphs[0]
    p_s8.text = "💡 핵심 설계: 모델을 맹신하는 구조가 아니라, 모델의 생성 결과를 서버가 엄격하게 검증하는 구조로 신뢰성을 확보했습니다."
    p_s8.font.name = FONT_FAMILY
    p_s8.font.size = Pt(11)
    p_s8.font.color.rgb = TEXT_WHITE
    p_s8.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 9 (NEW): Core Engineering Challenges & Solutions
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "핵심 기술 난제와 해결 과정", 
               category_text="VISIONINSIGHT  |  CORE ENGINEERING CHALLENGES & SOLUTIONS",
               subtitle_text="\"Gemini API 연결보다 어려웠던 것은, 검증 가능한 Agent를 만드는 과정\"")

    challenges_data = [
        ("HTTP 400", PRIMARY_CYAN, "구조화 출력 호환성", [
            ("문제: 복잡한 Schema 거부", "DiagnosisDraft 직접 전달 시 400 오류 (거부 속성 UNKNOWN)"),
            ("대응: 단순 JSON + Python 검증", "단순 JSON 전송 ➔ 서버에서 DiagnosisDraft 변환 및 Evidence 검증"),
            ("결과: 실제 T01/T03 Gemini E2E PASS", "단순 JSON 전환 후 실서버 통신 완전 성공 (PASS)"),
            ("교훈", "\"LLM은 생성, Python은 결정적 검증을 담당\"")
        ]),
        ("HTTP 504", WARNING_AMBER, "최종 보고서 생성 지연", [
            ("문제: 최종 JSON 요청 지연", "T03 도구 성공 후 최종 요청 14.15초에 504 발생 (원인 UNKNOWN)"),
            ("대응: 단계별 Timeout 계측", "조사(15초)와 생성 부하 분리 계측, 전체 60초 예산 유지"),
            ("결과: T03 단일 E2E PASS", "최종 30초 적용 시 18.967초(HTTP 200), 전체 24.600초 PASS"),
            ("교훈", "\"Timeout은 추측하지 않고 요청별로 계측한다\"")
        ]),
        ("T05 ERROR Evidence", DANGER_CORAL, "근거 검증 공백 극복", [
            ("문제: 필수 로그·미측정 근거 누락 위험", "오류 로그 미조회 완결 및 SKIPPED 모델 장애 오해 위험"),
            ("대응: ERROR 필수 17개 사실 검증", "로그 도구 강제, null 임의 변환 금지, 장애 가설 생성 차단"),
            ("결과: Mock PASS / 실서버 NOT TESTED", "오프라인 886개 회귀 테스트 통과 (실제 Gemini E2E 미호출)"),
            ("교훈", "\"유효한 JSON이 반드시 근거 있는 진단 보고서는 아니다\"")
        ])
    ]

    card_w9 = Inches(3.7)
    cx9 = 0.8
    for ctitle, ccol, csub, citems in challenges_data:
        add_card(s9, Inches(cx9), Inches(1.4), card_w9, Inches(4.35), border_color=ccol)
        tb_c = s9.shapes.add_textbox(Inches(cx9 + 0.18), Inches(1.50), card_w9 - Inches(0.36), Inches(4.15))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        
        p = tf_c.paragraphs[0]
        p.text = ctitle
        p.font.name = FONT_FAMILY
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ccol

        p_sub = tf_c.add_paragraph()
        p_sub.text = "• " + csub
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(10.5)
        p_sub.font.bold = True
        p_sub.font.color.rgb = TEXT_WHITE
        p_sub.space_before = Pt(2)

        for label, desc in citems:
            p_l = tf_c.add_paragraph()
            p_l.text = f"• {label}"
            p_l.font.name = FONT_FAMILY
            p_l.font.size = Pt(9.5)
            p_l.font.bold = True
            p_l.font.color.rgb = PRIMARY_CYAN if "교훈" in label else SUCCESS_GREEN if "결과" in label else WARNING_AMBER
            p_l.space_before = Pt(5)

            p_d = tf_c.add_paragraph()
            p_d.text = "   " + desc
            p_d.font.name = FONT_FAMILY
            p_d.font.size = Pt(9.0)
            p_d.font.color.rgb = TEXT_LIGHT

        cx9 += 4.0

    add_card(s9, Inches(0.8), Inches(5.95), Inches(11.733), Inches(0.8), bg_color=CARD_BG_ALT, border_color=PRIMARY_CYAN)
    tb_sum9 = s9.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(11.3), Inches(0.7))
    tf_sum9 = tb_sum9.text_frame
    p_sum9 = tf_sum9.paragraphs[0]
    p_sum9.text = "💡 핵심 메시지: AI가 그럴듯한 답을 만드는 것보다, 서버가 그 답을 검증하고 근거 부족 시 거부하는 구조가 중요했습니다."
    p_sum9.font.name = FONT_FAMILY
    p_sum9.font.size = Pt(11)
    p_sum9.font.bold = True
    p_sum9.font.color.rgb = TEXT_WHITE
    p_sum9.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 10: Real Server E2E: T01 & T03 (Former Slide 9)
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "실제 Gemini API E2E 검증 결과: T01 정상 및 T03 경고 검증 (PASS)")

    # Left: T01
    add_card(s10, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.3), border_color=SUCCESS_GREEN)
    tb_t01 = s10.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_01 = tb_t01.text_frame
    tf_01.word_wrap = True
    p = tf_01.paragraphs[0]
    p.text = "T01 정상 시나리오: 실서버 E2E 성공 (PASS)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    t01_bullets_final = [
        ("검사 식별자", "INSP-001396 (정상 병 검사)"),
        ("판정 결과", "모델 OK  /  서버 최종 OK (ALL_CHECKS_PASSED)"),
        ("FastAPI 응답", "HTTP 200 OK (analysis_status: COMPLETED)"),
        ("Gemini 통신", "2회 모두 HTTP 200 (전체 소요: 13.578초)"),
        ("토큰 사용량", "총 4,706 토큰 (입력 3,115, 출력 125, Thinking 1,466)"),
        ("도구 실행", "get_inspection_result 실행 (2.23ms, E0001)"),
        ("근거 검증", "관측 사실 6개 (F0001~F0006) 1:1 결정적 검증 통과")
    ]
    for k, v in t01_bullets_final:
        p = tf_01.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(7)

    # Right: T03
    add_card(s10, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.3), border_color=SUCCESS_GREEN)
    tb_t03 = s10.shapes.add_textbox(Inches(7.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_03 = tb_t03.text_frame
    tf_03.word_wrap = True
    p = tf_03.paragraphs[0]
    p.text = "T03 조도 저하 시나리오: 실서버 E2E 성공 (PASS)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    t03_bullets_final = [
        ("검사 식별자", "INSP-001398 (조도 저하 주입 검사)"),
        ("판정 결과", "모델 OK  /  서버 최종 REINSPECT (IMAGE_QUALITY_WARNING)"),
        ("FastAPI 응답", "HTTP 200 OK (analysis_status: COMPLETED, 회차 final30_1)"),
        ("Gemini 통신", "3회 모두 HTTP 200 (전체 소요: 24.600초)"),
        ("토큰 사용량", "총 9,873 토큰 (입력 5,576, 출력 345, Thinking 3,952)"),
        ("도구 실행 (2개)", "get_inspection_result(E0001) + analyze_image_quality(E0002)"),
        ("근거 검증", "관측 사실 12개 (F0001~F0012) 100% 일치 검증 통과"),
        ("요청별 실측치", "도구1: 2.43s (200) | 도구2: 3.05s (200) | 최종: 18.967s (200)")
    ]
    for k, v in t03_bullets_final:
        p = tf_03.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.2)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(6)

    p_note9 = tf_03.add_paragraph()
    p_note9.text = "※ 타임아웃 주석: 최종 JSON 요청에 한해 30초 timeout을 선택 적용하여 18.967초에 성공. 과거 504 원인은 UNKNOWN이며 단일 회차 성공으로 향후 안정성을 보장하지 않음."
    p_note9.font.name = FONT_FAMILY
    p_note9.font.size = Pt(9.0)
    p_note9.font.color.rgb = WARNING_AMBER
    p_note9.space_before = Pt(6)

    # ==========================================
    # SLIDE 11: Real Demo & Presentation Flow (Former Slide 10)
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "발표 데모 시연: T01 양품 vs T03 조도 경고 정적 데모")

    add_card(s11, Inches(0.8), Inches(1.4), Inches(4.5), Inches(5.3))
    tb_d10 = s11.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(3.9), Inches(4.8))
    tf_d10 = tb_d10.text_frame
    tf_d10.word_wrap = True
    p = tf_d10.paragraphs[0]
    p.text = "발표 시연 흐름 제안"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_CYAN

    d_steps_final = [
        ("Step 1. T01 정상 검사 데모", "t01-report-demo.html 시연 (모델 OK ➔ 최종 OK, 도구 1개, 팩트 6개)"),
        ("Step 2. T03 조도 경고 데모", "t03-report-demo.html 시연 (평균 밝기 81.24 저하 감지)"),
        ("Step 3. 판정 차이 설명", "두 검사 모두 모델 판정은 OK지만, T03은 Rule 2로 REINSPECT 분류"),
        ("Step 4. 판정 불변 원칙", "Gemini는 판정을 변경하지 않고 도구 2개로 팩트 12개 기반 사후 진단 보조")
    ]
    for stitle, sdesc in d_steps_final:
        p = tf_d10.add_paragraph()
        p.text = stitle
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)
        p_d = tf_d10.add_paragraph()
        p_d.text = "  " + sdesc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_MUTED

    p_warn10 = tf_d10.add_paragraph()
    p_warn10.text = "※ 정적 데모 원칙: 발표 현장 네트워크 장애 및 과금을 방지하기 위해 실서버 검증 데이터를 재생하는 독립형 정적 데모로 시연함."
    p_warn10.font.name = FONT_FAMILY
    p_warn10.font.size = Pt(9.2)
    p_warn10.font.color.rgb = WARNING_AMBER
    p_warn10.space_before = Pt(14)

    add_card(s11, Inches(5.6), Inches(1.4), Inches(6.933), Inches(5.3), border_color=PRIMARY_CYAN)
    tb_j10 = s11.shapes.add_textbox(Inches(5.9), Inches(1.65), Inches(6.3), Inches(4.8))
    tf_j10 = tb_j10.text_frame
    tf_j10.word_wrap = True
    p = tf_j10.paragraphs[0]
    p.text = "실서버 성공 진단 보고서 비교 (T01 vs T03)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    json_highlights_final = [
        ("T01 종합 요약", "\"기록된 모델 판정과 종합 판정은 OK입니다. 실제 무결점 제품을 보증하지 않습니다.\"", TEXT_LIGHT),
        ("T01 권장 조치", "check_code: CONTINUE_MONITORING (evidence_refs: [\"F0003\"])", SUCCESS_GREEN),
        ("T01 관측 사실", "F0001~F0006 (6개, get_inspection_result E0001 매핑)", PRIMARY_CYAN),
        ("T03 종합 요약", "\"기록된 종합 판정은 REINSPECT입니다. 제품 불량 확정과 구분하여 재검사가 필요합니다.\"", WARNING_AMBER),
        ("T03 권장 조치", "LIGHTING_EXPOSURE_CHECK ([\"F0004\",\"F0007\"]) / MANUAL_REINSPECTION ([\"F0003\"])", WARNING_AMBER),
        ("T03 관측 사실", "F0001~F0012 (12개, 결과 E0001 + 품질 E0002 매핑, 100% 일치)", SUCCESS_GREEN),
        ("공통 안전 한계", "limitations 4건 필수 고지 (물리적 고장 미확정, 모델 비보증, 기준치 미기록 등)", TEXT_MUTED)
    ]
    for k, v, c in json_highlights_final:
        p = tf_j10.add_paragraph()
        p.text = f"• {k}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = c
        p.space_before = Pt(5)
        p_sub = tf_j10.add_paragraph()
        p_sub.text = f"   {v}"
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(9.5)
        p_sub.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 12: Implementation Status (Former Slide 11)
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "구현 현황 및 향후 계획 (투명한 성과 구분)")

    col_width = Inches(3.7)
    cols_data_final = [
        ("1. 구현 및 실서버 PASS",
         SUCCESS_GREEN,
         ["FastAPI 백엔드 & 고유 ID 발급",
          "PatchCore CPU 추론 엔진 (42.53ms)",
          "규칙 기반 종합 판정 (REINSPECT 분리)",
          "Gemini 4대 조회 도구 바인딩",
          "결정적 근거 검증 (Evidence/Pointer)",
          "단순 JSON 스키마 호환 구조",
          "T01 실서버 Gemini E2E (PASS)",
          "T03 실서버 Gemini E2E (PASS)"]),

        ("2. 정책 보완 & Mock PASS",
         PRIMARY_CYAN,
         ["전체 pytest 886개 PASS (경고 2건)",
          "T03 경고 조기 종료 오프라인 검증",
          "T05 ERROR 운영 근거 정책 보완 완료",
          "  - 필수 17개 사실 강제 (로그/미측정)",
          "  - SKIPPED 모델 장애 가설 차단",
          "T05 계측기 정정 완료 (실행 준비: READY)",
          "※ T05 실제 Gemini E2E: NOT TESTED",
          "※ T06 미탐 사례 E2E: NOT TESTED"]),

        ("3. 향후 개발 예정",
         WARNING_AMBER,
         ["React 웹 관제 대시보드 (Day 6)",
          "컨베이어 주행 동영상 Live View",
          "실시간 프레임 Grab & 마우스 ROI",
          "ROI 크롭 이미지 PatchCore 검사",
          "Google Cloud Run 컨테이너 배포 (Day 7)",
          "T05/T06 단일 실서버 E2E 검증"])
    ]
    cx = 0.8
    for ctitle, ccol, citems in cols_data_final:
        add_card(s12, Inches(cx), Inches(1.4), col_width, Inches(5.3), border_color=ccol)
        tb_col = s12.shapes.add_textbox(Inches(cx + 0.2), Inches(1.65), col_width - Inches(0.4), Inches(4.8))
        tf_col = tb_col.text_frame
        tf_col.word_wrap = True
        p = tf_col.paragraphs[0]
        p.text = ctitle
        p.font.name = FONT_FAMILY
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ccol

        for item in citems:
            p_i = tf_col.add_paragraph()
            p_i.text = "• " + item if not item.startswith("  -") and not item.startswith("※") else item
            p_i.font.name = FONT_FAMILY
            p_i.font.size = Pt(10.2)
            p_i.font.color.rgb = WARNING_AMBER if item.startswith("※") else TEXT_LIGHT
            p_i.space_before = Pt(6)

        cx += 4.0

    # ==========================================
    # SLIDE 13: Limitations & Future Work (Former Slide 12)
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13)
    add_header(s13, "한계 및 향후 개선 계획: 솔직한 엔지니어링 고찰")

    limits_final = [
        ("1. PatchCore 미탐 한계 (Recall 50.8%)",
         DANGER_CORAL,
         "현재 ResNet-18 모델은 83장 테스트셋에서 미세 파손/오염 31건을 미검출함 (T06 합격 사례 존재). 모델 OK가 제품 무결점을 보증하지 않음.",
         "개선 계획: WideResNet-50 등 백본 교체나 다해상도 패치 결합을 개선 후보로 검토 (성능 보장 아님)"),

        ("2. Gemini HTTP 504 및 타임아웃 안정성 이슈",
         WARNING_AMBER,
         "과거 T03 최종 JSON 요청에서 HTTP 504가 발생했으며 정확한 서버 내부 원인은 UNKNOWN입니다. 이후 최종 JSON 요청에만 30초 timeout을 적용한 단일 재검증에서 해당 요청은 18.967초 만에 HTTP 200으로 완료됐고, 전체 E2E는 24.600초에 PASS했습니다. 기본 timeout은 15초이며, 단일 성공 실측치로 향후 반복 안정성을 보증하지는 않습니다.",
         "개선 계획: 프롬프트 압축, 비동기 큐 전환, 반복 안정성 계측 및 운영 모니터링 체계 확립 필요"),

        ("3. 외부 LLM 호출 지연 (T01: 13.58초, T03: 24.60초)",
         PRIMARY_CYAN,
         "Gemini API 호출에 수십 초가 소요되어 생산 라인 초고속 동기 인라인 제어에는 부적합하며, 실제 PLC/배출기 제어는 수행하지 않음.",
         "개선 계획: 사후 비동기 진단 워크플로우로 한정 운영하고, 라인 제어는 서버의 실시간 규칙 판정값에만 전담")
    ]
    ly = 1.4
    for ltitle, lcol, lprob, lsol in limits_final:
        add_card(s13, Inches(0.8), Inches(ly), Inches(11.733), Inches(1.6), border_color=lcol)
        tb_l = s13.shapes.add_textbox(Inches(1.1), Inches(ly + 0.12), Inches(11.1), Inches(1.4))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True
        p = tf_l.paragraphs[0]
        p.text = ltitle
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = lcol

        p_p = tf_l.add_paragraph()
        p_p.text = "• 한계: " + lprob
        p_p.font.name = FONT_FAMILY
        p_p.font.size = Pt(10.2)
        p_p.font.color.rgb = TEXT_LIGHT
        p_p.space_before = Pt(3)

        p_s = tf_l.add_paragraph()
        p_s.text = "✔ 검토: " + lsol
        p_s.font.name = FONT_FAMILY
        p_s.font.size = Pt(10.2)
        p_s.font.color.rgb = SUCCESS_GREEN
        p_s.space_before = Pt(2)

        ly += 1.8

    # ==========================================
    # SLIDE 14: Conclusion (Former Slide 13)
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14)

    add_card(s14, Inches(0.8), Inches(1.1), Inches(11.733), Inches(5.4), bg_color=CARD_BG_ALT, border_color=PRIMARY_CYAN)

    tb_end = s14.shapes.add_textbox(Inches(1.3), Inches(1.5), Inches(10.7), Inches(4.5))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True

    p = tf_end.paragraphs[0]
    p.text = "VisionInsight — 근거 기반 산업 비전 AI Agent 설계"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_CYAN

    conclusions_final = [
        ("PatchCore", "이상 점수 산출 및 모델 판정 소견(result) 제공 (CPU 42.53ms)"),
        ("FastAPI", "서버 규칙 기반 최종 검사 판정(final_result: REINSPECT 분리)과 감사 로그 기록"),
        ("Gemini", "도구 조회를 통한 근거 기반 진단과 점검 권고 제안 (사후 진단 보조)"),
        ("Python", "엄격한 근거 검증(Evidence/Pointer)과 판정 불변 보호 정책으로 시스템 신뢰성 확보")
    ]
    for ctitle, cdesc in conclusions_final:
        p_c = tf_end.add_paragraph()
        p_c.text = f"✔ {ctitle}: {cdesc}"
        p_c.font.name = FONT_FAMILY
        p_c.font.size = Pt(12)
        p_c.font.bold = True
        p_c.font.color.rgb = TEXT_LIGHT
        p_c.space_before = Pt(12)

    p_msg = tf_end.add_paragraph()
    p_msg.text = "\"AI가 검사 판정을 대신하는 것이 아니라, 검증된 검사 정보를 운영자가 이해하고 활용하도록 돕는 시스템\""
    p_msg.font.name = FONT_FAMILY
    p_msg.font.size = Pt(13)
    p_msg.font.bold = True
    p_msg.font.color.rgb = PRIMARY_CYAN
    p_msg.space_before = Pt(16)
    p_msg.alignment = PP_ALIGN.CENTER

    p_qa = tf_end.add_paragraph()
    p_qa.text = "경청해 주셔서 감사합니다.  |  Q&A (질의응답)"
    p_qa.font.name = FONT_FAMILY
    p_qa.font.size = Pt(18)
    p_qa.font.bold = True
    p_qa.font.color.rgb = SUCCESS_GREEN
    p_qa.space_before = Pt(16)
    p_qa.alignment = PP_ALIGN.CENTER

    # Save 14-slide PPTX
    output_path = os.path.join(os.path.dirname(__file__), "VisionInsight_Presentation_14slides.pptx")
    prs.save(output_path)
    print(f"VisionInsight Presentation 14slides successfully created at: {output_path}")

if __name__ == "__main__":
    create_presentation_14slides()
