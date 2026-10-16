"""VisionInsight 프레젠테이션 v3 생성 스크립트
기술 내용 검수 및 실측 사실 기반 정밀 교정 (13장 슬라이드 완전 생성)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation_v3():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette (Industrial Dark Navy & Modern Accents)
    BG_COLOR = RGBColor(12, 19, 34)          # #0C1322
    CARD_BG = RGBColor(22, 30, 46)           # #161E2E
    CARD_BG_ALT = RGBColor(16, 23, 38)       # #101726
    CARD_BORDER = RGBColor(42, 54, 73)       # #2A3649
    PRIMARY_CYAN = RGBColor(76, 215, 246)    # #4CD7F6
    SUCCESS_GREEN = RGBColor(52, 211, 153)   # #34D399
    WARNING_AMBER = RGBColor(251, 191, 36)   # #FBBF24
    DANGER_CORAL = RGBColor(248, 113, 113)   # #F87171
    TEXT_WHITE = RGBColor(248, 250, 252)     # #F8FAFC
    TEXT_LIGHT = RGBColor(203, 213, 225)     # #CBD5E1
    TEXT_MUTED = RGBColor(148, 163, 184)     # #94A3B8

    FONT_FAMILY = "Malgun Gothic"

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR
        return bg

    def add_header(slide, title_text, category_text="VISIONINSIGHT  |  INDUSTRIAL VISION AI AGENT"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.28))
        tf = cat_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = category_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_CYAN

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.7), Inches(0.55))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_top = tf_t.margin_bottom = tf_t.margin_left = tf_t.margin_right = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(21)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    # ==========================================
    # SLIDE 1: Title Slide (v3 Verified)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    add_card(s1, Inches(0.8), Inches(1.1), Inches(11.733), Inches(5.4), bg_color=CARD_BG_ALT, border_color=PRIMARY_CYAN)

    tbox = s1.shapes.add_textbox(Inches(1.4), Inches(1.7), Inches(10.5), Inches(2.2))
    tf = tbox.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "VisionInsight"
    p0.font.name = FONT_FAMILY
    p0.font.size = Pt(46)
    p0.font.bold = True
    p0.font.color.rgb = PRIMARY_CYAN

    p1 = tf.add_paragraph()
    p1.text = "AI 기반 병 외관 이상 탐지 및 근거 기반 진단 시스템"
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_before = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "MVTec Bottle 기반 PatchCore 이상 탐지와 Gemini Function Calling의 결합"
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(13)
    p2.font.color.rgb = TEXT_LIGHT
    p2.space_before = Pt(12)

    badges = [
        ("PatchCore 비지도 학습", PRIMARY_CYAN),
        ("개발 환경 추론 42.5ms", SUCCESS_GREEN),
        ("대표 판정 규칙 4가지", WARNING_AMBER),
        ("Gemini 도구 자율 조사", PRIMARY_CYAN)
    ]
    bx = 1.4
    for btext, bcol in badges:
        add_card(s1, Inches(bx), Inches(4.3), Inches(2.4), Inches(0.5), bg_color=CARD_BG, border_color=bcol)
        tb = s1.shapes.add_textbox(Inches(bx), Inches(4.3), Inches(2.4), Inches(0.5))
        tf_b = tb.text_frame
        tf_b.margin_top = Inches(0.1)
        p = tf_b.paragraphs[0]
        p.text = btext
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = bcol
        p.alignment = PP_ALIGN.CENTER
        bx += 2.6

    fbox = s1.shapes.add_textbox(Inches(1.4), Inches(5.35), Inches(10.5), Inches(0.5))
    tf_f = fbox.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "발표 일자: 2026.10.20  |  VisionInsight 개발 프로젝트"
    p_f.font.name = FONT_FAMILY
    p_f.font.size = Pt(11)
    p_f.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 2: Problem Definition (v3 Verified)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "문제 정의: 인라인 비전 검사의 한계와 해결 접근법")

    add_card(s2, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.3), border_color=DANGER_CORAL)
    tb_l = s2.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "기존 비전 검사의 주요 한계"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = DANGER_CORAL

    points_l = [
        ("블랙박스 이상 점수 (설명 부재)", "이상 점수만 산출되어 결함의 원인과 현장 조치 방법을 파악하기 어려움"),
        ("환경 변화로 인한 오판 위험", "조도 저하 등 촬영 환경 이상을 제품 결함으로 오진단할 위험 존재"),
        ("운영자 분석 부담", "이상 발생 시 로그 확인과 설비 점검을 위해 작업자의 추가 검토 필요")
    ]
    for title, desc in points_l:
        p = tf_l.add_paragraph()
        p.text = "• " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(16)
        p_sub = tf_l.add_paragraph()
        p_sub.text = "   " + desc
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = TEXT_LIGHT
        p_sub.space_before = Pt(3)

    add_card(s2, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.3), border_color=SUCCESS_GREEN)
    tb_r = s2.shapes.add_textbox(Inches(7.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "VisionInsight의 해결 접근법"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    points_r = [
        ("비지도 이상 탐지 (PatchCore)", "양품 180장 Coreset 학습 기반 이상 점수 산출 (개발 환경 42.5ms)"),
        ("품질 이상과 결함의 판정 분리", "조도 이상을 제품 결함과 구별하여 재검사(REINSPECT) 대상으로 분류하도록 설계"),
        ("자율 조사 Gemini AI Agent", "4대 도구로 팩트를 조회하여 원인 조사와 점검 의사결정 지원")
    ]
    for title, desc in points_r:
        p = tf_r.add_paragraph()
        p.text = "✔ " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_CYAN
        p.space_before = Pt(16)
        p_sub = tf_r.add_paragraph()
        p_sub.text = "   " + desc
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = TEXT_LIGHT
        p_sub.space_before = Pt(3)

    # ==========================================
    # SLIDE 3: System Architecture (v3 Verified)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "시스템 아키텍처: 3단계 책임 분리 (Vision — Decision — Agent)")

    layers_v3 = [
        ("Layer 1: Inspection Engine (Vision & ML)",
         "OpenCV 조도 분석 + PatchCore CPU 추론",
         PRIMARY_CYAN,
         "OpenCV 밝기 통계(mean/std) 산출  ➔  ResNet-18 PatchCore 추론 (개발 환경 42.5ms)  ➔  이상 점수 및 모델 소견(result: OK/NG/UNKNOWN)"),

        ("Layer 2: Deterministic Decision Engine (Business Logic)",
         "서버 규칙 기반 최종 검사 판정 + 감사 로그 저장",
         SUCCESS_GREEN,
         "6대 우선순위 규칙 집행  ➔  서버 규칙 기반 final_result(OK/NG/REINSPECT/ERROR) 확정  ➔  JSONL 감사 로그 영구화"),

        ("Layer 3: Gemini AI Agent Service (Post-Diagnosis)",
         "도구 호출 제안 + 서버 실행 + 결정적 근거 검증",
         WARNING_AMBER,
         "Gemini가 조회 제안  ➔  서버가 4대 도구 안전 실행  ➔  Evidence ID / JSON Pointer 대조  ➔  근거 기반 사후 진단")
    ]

    ty = 1.45
    for ltitle, lsub, lcol, lbody in layers_v3:
        add_card(s3, Inches(0.8), Inches(ty), Inches(11.733), Inches(1.4), bg_color=CARD_BG, border_color=lcol)
        tb = s3.shapes.add_textbox(Inches(1.1), Inches(ty + 0.15), Inches(11.1), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = ltitle + "  |  " + lsub
        p.font.name = FONT_FAMILY
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = lcol

        p_b = tf.add_paragraph()
        p_b.text = lbody
        p_b.font.name = FONT_FAMILY
        p_b.font.size = Pt(11.5)
        p_b.font.color.rgb = TEXT_LIGHT
        p_b.space_before = Pt(8)

        ty += 1.6

    add_card(s3, Inches(0.8), Inches(6.35), Inches(11.733), Inches(0.55), bg_color=CARD_BG_ALT, border_color=PRIMARY_CYAN)
    tb_call = s3.shapes.add_textbox(Inches(1.0), Inches(6.35), Inches(11.3), Inches(0.55))
    tf_c = tb_call.text_frame
    p_c = tf_c.paragraphs[0]
    p_c.text = "🛡 핵심 안전 정책: AI Agent는 설명과 점검 권고를 담당하며, 서버의 원본 검사 판정(final_result)을 수정할 수 없습니다."
    p_c.font.name = FONT_FAMILY
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = TEXT_WHITE
    p_c.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 4: Dataset & PatchCore Evaluation (v3 Verified)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "비지도 이상 탐지: PatchCore 엔진 및 실측 평가 성능")

    add_card(s4, Inches(0.8), Inches(1.4), Inches(4.5), Inches(5.3))
    tb_s4 = s4.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(3.9), Inches(4.8))
    tf_s4 = tb_s4.text_frame
    tf_s4.word_wrap = True
    p = tf_s4.paragraphs[0]
    p.text = "PatchCore 모델 사양"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_CYAN

    s4_specs = [
        ("학습 데이터", "MVTec Bottle 정상 180장 Coreset"),
        ("백본 모델", "ResNet-18 (ImageNet 사전학습 가중치)"),
        ("메모리 뱅크", "형상 [9216, 384], 크기 24.18MB"),
        ("결정 임계값", "독립 정상 29장 통계 기반 34.1442 (μ+3σ)"),
        ("추론 지연시간", "개발 환경(x86_64 CPU) 단일 이미지 평균 42.53ms")
    ]
    for k, v in s4_specs:
        p = tf_s4.add_paragraph()
        p.text = f"• {k}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)
        p_sub = tf_s4.add_paragraph()
        p_sub.text = f"   {v}"
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(10.5)
        p_sub.font.color.rgb = TEXT_MUTED

    stat_cards_v3 = [
        ("Image AUROC", "0.8286", "이상 점수의 양품/불량 분별력", PRIMARY_CYAN, Inches(5.6), Inches(1.4)),
        ("정밀도 (Precision)", "100.0%", "근거: TP=32, FP=0 (해당 83장 표본 기준)", SUCCESS_GREEN, Inches(9.3), Inches(1.4)),
        ("재현율 (Recall)", "50.79%", "근거: TP=32, FN=31 (미탐 31건 발생)", WARNING_AMBER, Inches(5.6), Inches(3.4)),
        ("CPU 추론 시간", "42.53 ms", "개발 환경 단일 이미지 평균 실측치", SUCCESS_GREEN, Inches(9.3), Inches(3.4))
    ]
    for title, val, desc, col, l, t in stat_cards_v3:
        add_card(s4, l, t, Inches(3.45), Inches(1.75), border_color=col)
        tb_stat = s4.shapes.add_textbox(l + Inches(0.2), t + Inches(0.15), Inches(3.05), Inches(1.45))
        tf_stat = tb_stat.text_frame
        tf_stat.word_wrap = True
        p = tf_stat.paragraphs[0]
        p.text = title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED

        p_v = tf_stat.add_paragraph()
        p_v.text = val
        p_v.font.name = FONT_FAMILY
        p_v.font.size = Pt(28)
        p_v.font.bold = True
        p_v.font.color.rgb = col
        p_v.space_before = Pt(2)

        p_d = tf_stat.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_LIGHT

    add_card(s4, Inches(5.6), Inches(5.4), Inches(7.15), Inches(1.3), bg_color=CARD_BG_ALT, border_color=DANGER_CORAL)
    tb_disc = s4.shapes.add_textbox(Inches(5.8), Inches(5.45), Inches(6.75), Inches(1.2))
    tf_d = tb_disc.text_frame
    tf_d.word_wrap = True
    p = tf_d.paragraphs[0]
    p.text = "⚠ 테스트 표본 한계 및 미탐 사례 안내"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = DANGER_CORAL
    p_sub = tf_d.add_paragraph()
    p_sub.text = "• 독립 테스트셋 83장: 정상 20장(TN 20, FP 0), 불량 63장(TP 32, FN 31) ➔ 정밀도 100%는 실전 보장이 아님\n• T06(실제 파손 불량)은 이상 점수 32.27로 합격 처리됨 ➔ 미탐 한계를 인위적 조작 없이 보존 (정답 라벨은 Agent에 미제공)"
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(9.5)
    p_sub.font.color.rgb = TEXT_LIGHT
    p_sub.space_before = Pt(3)

    # ==========================================
    # SLIDE 5: Inspection Pipeline & Rules (v3 B안 대표 4규칙)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "검사 판정 체계: 대표 판정 규칙 4가지 (6대 우선순위 계층 기반)")

    add_card(s5, Inches(0.8), Inches(1.4), Inches(11.733), Inches(1.3), border_color=PRIMARY_CYAN)
    tb_c5 = s5.shapes.add_textbox(Inches(1.1), Inches(1.5), Inches(11.1), Inches(1.1))
    tf_c5 = tb_c5.text_frame
    tf_c5.word_wrap = True
    p = tf_c5.paragraphs[0]
    p.text = "핵심 원칙: 모델 소견(result)과 최종 검사 판정(final_result)의 명확한 분리"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_CYAN

    p1 = tf_c5.add_paragraph()
    p1.text = "• result (AI 모델 소견): PatchCore가 외관을 평가한 결과 (OK / NG / UNKNOWN)"
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(11)
    p1.font.color.rgb = TEXT_WHITE
    p1.space_before = Pt(4)

    p2 = tf_c5.add_paragraph()
    p2.text = "• final_result (서버 종합 판정): 영상 품질과 시스템 상태를 반영한 최종 판정 (OK / NG / REINSPECT / ERROR)"
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = SUCCESS_GREEN
    p2.space_before = Pt(2)

    rule_boxes_v3 = [
        ("Rule 1: 파이프라인 실패", "ERROR", DANGER_CORAL, "INSPECTION_FAILED", "status == 'FAILED' 또는 이미지 로드 실패 시"),
        ("Rule 2: 영상 품질 경고", "REINSPECT", WARNING_AMBER, "IMAGE_QUALITY_WARNING", "image_quality_status != 'NORMAL' (조도 저하/과다 시 오판 방지 재검)"),
        ("Rule 4: 제품 결함 검출", "NG", DANGER_CORAL, "PRODUCT_ANOMALY_DETECTED", "품질 정상 상태에서 result == 'NG' (이상 점수 > 34.1442)"),
        ("Rule 5: 정상 통과", "OK", SUCCESS_GREEN, "ALL_CHECKS_PASSED", "품질 정상 상태에서 result == 'OK' (이상 점수 <= 34.1442)")
    ]
    r_coords = [
        (Inches(0.8), Inches(2.95)), (Inches(6.8), Inches(2.95)),
        (Inches(0.8), Inches(4.75)), (Inches(6.8), Inches(4.75))
    ]
    for (rtitle, rres, rcol, rcode, rdesc), (rx, ry) in zip(rule_boxes_v3, r_coords):
        add_card(s5, rx, ry, Inches(5.7), Inches(1.55), border_color=rcol)
        tb_r = s5.shapes.add_textbox(rx + Inches(0.2), ry + Inches(0.15), Inches(5.3), Inches(1.25))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True
        p = tf_r.paragraphs[0]
        p.text = f"{rtitle}  ➔  final_result: {rres}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = rcol

        p_c = tf_r.add_paragraph()
        p_c.text = f"사유 코드: {rcode}"
        p_c.font.name = FONT_FAMILY
        p_c.font.size = Pt(10.5)
        p_c.font.color.rgb = TEXT_MUTED
        p_c.space_before = Pt(2)

        p_d = tf_r.add_paragraph()
        p_d.text = rdesc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_before = Pt(4)

    # ==========================================
    # SLIDE 6: Gemini Function Calling (v3 Clarified)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Gemini AI Agent와 Function Calling: 자율 조사 4대 도구")

    add_card(s6, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.3), border_color=PRIMARY_CYAN)
    tb_s6_l = s6.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_s6_l = tb_s6_l.text_frame
    tf_s6_l.word_wrap = True
    p = tf_s6_l.paragraphs[0]
    p.text = "도구 호출 제어 흐름 (서버 안전 실행)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_CYAN

    flow_steps_v3 = [
        ("1. Gemini Function Call 제안", "Gemini가 조회할 도구명과 인자를 JSON으로 선언 (SDK 자동 실행 차단)"),
        ("2. FastAPI AgentService 검증", "서버가 호출 인자 유효성과 데이터 격리 정책 검증"),
        ("3. Python 도구 안전 실행", "서버가 로컬 Python 함수를 대신 실행하고 Evidence ID 부여"),
        ("4. FunctionResponse 전달", "도구 실행 결과를 role='user'로 Gemini에 전달하여 팩트 기반 보고서 유도")
    ]
    for step_title, step_desc in flow_steps_v3:
        p = tf_s6_l.add_paragraph()
        p.text = step_title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)
        p_d = tf_s6_l.add_paragraph()
        p_d.text = "  ↳ " + step_desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT

    add_card(s6, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.3), border_color=SUCCESS_GREEN)
    tb_s6_r = s6.shapes.add_textbox(Inches(7.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_s6_r = tb_s6_r.text_frame
    tf_s6_r.word_wrap = True
    p = tf_s6_r.paragraphs[0]
    p.text = "4대 Python 조회 도구 (InspectionTools)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    tools_v3 = [
        ("get_inspection_result", "검사 상태, 이상 점수, 임계값, 판정 결과 조회"),
        ("analyze_image_quality", "OpenCV 측정 밝기 통계(mean/std) 및 기준 통계 조회"),
        ("get_execution_logs", "시작~종료 단계별 라이프사이클 이벤트 감사 로그 조회"),
        ("get_inspection_history", "동일 제품군 과거 검사 이력 및 추세 조회")
    ]
    for t_name, t_desc in tools_v3:
        p = tf_s6_r.add_paragraph()
        p.text = "✔ " + t_name
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_CYAN
        p.space_before = Pt(10)
        p_d = tf_s6_r.add_paragraph()
        p_d.text = "   " + t_desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 7: Deterministic Grounding & Guardrails (v3 Verified)
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "근거 기반 검증과 판정 보호: 환각 억제 및 판정 보존")

    add_card(s7, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.3), border_color=PRIMARY_CYAN)
    tb_s7_l = s7.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_s7_l = tb_s7_l.text_frame
    tf_s7_l.word_wrap = True
    p = tf_s7_l.paragraphs[0]
    p.text = "결정적 근거 검증 (Deterministic Grounding)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_CYAN

    g_items_v3 = [
        ("Evidence ID 영수증 체계", "도구 성공 실행 시마다 고유 영수증 ID(E0001 등) 발급"),
        ("RFC 6901 JSON Pointer 참조", "데이터 위치(/data/status 등)를 정확한 포인터로 명시"),
        ("서버 결정적 값 비교", "JSON Pointer로 참조한 실제 도구 반환값과 모델의 관측값을 결정적으로 비교"),
        ("불일치 시 보고서 거부", "검증 대상 필드 불일치 시 EVIDENCE_MISMATCH (HTTP 502)로 즉시 차단")
    ]
    for title, desc in g_items_v3:
        p = tf_s7_l.add_paragraph()
        p.text = "• " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(12)
        p_d = tf_s7_l.add_paragraph()
        p_d.text = "   " + desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_before = Pt(2)

    add_card(s7, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.3), border_color=WARNING_AMBER)
    tb_s7_r = s7.shapes.add_textbox(Inches(7.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_s7_r = tb_s7_r.text_frame
    tf_s7_r.word_wrap = True
    p = tf_s7_r.paragraphs[0]
    p.text = "판정 보호 정책 (Decision Guardrail)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = WARNING_AMBER

    gr_items_v3 = [
        ("원본 검사 판정 보존", "Agent 응답의 판정 필드는 검사 파이프라인의 원래 값 일치 여부를 검사"),
        ("판정 임의 수정 차단", "Gemini가 서버의 final_result를 임의로 번복하거나 수정할 수 없음"),
        ("가설과 관측의 분리", "조도 저하가 관측되어도 조명 고장으로 단정하지 않고 점검 권고로 분리"),
        ("데이터 격리 및 별칭화", "불투명 별칭(TARGET-PRODUCT)과 중립 ID로 원본 정답 라벨 완전 은닉")
    ]
    for title, desc in gr_items_v3:
        p = tf_s7_r.add_paragraph()
        p.text = "✔ " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(12)
        p_d = tf_s7_r.add_paragraph()
        p_d.text = "   " + desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_before = Pt(2)

    # ==========================================
    # SLIDE 8: Gemini Structured Output (v3 Verified)
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "기술적 문제 해결: Gemini 구조화 출력 호환성 개선")

    add_card(s8, Inches(0.8), Inches(1.4), Inches(5.7), Inches(4.3), border_color=DANGER_CORAL)
    tb_b8 = s8.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(5.1), Inches(3.8))
    tf_b8 = tb_b8.text_frame
    tf_b8.word_wrap = True
    p = tf_b8.paragraphs[0]
    p.text = "Before: 직접 구조화 출력 시도 ➔ HTTP 400"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = DANGER_CORAL

    b_points_v3 = [
        ("복잡한 Pydantic 스키마 전달", "DiagnosisDraft 스키마를 response_json_schema로 직접 전송"),
        ("실서버 HTTP 400 거부", "Gemini 서버가 스키마 호환성 문제로 ClientError 반환 (정확한 거부 속성은 UNKNOWN)"),
        ("결과", "사후 진단 보고서 생성 실패")
    ]
    for title, desc in b_points_v3:
        p = tf_b8.add_paragraph()
        p.text = "• " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)
        p_d = tf_b8.add_paragraph()
        p_d.text = "   " + desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT

    add_card(s8, Inches(6.8), Inches(1.4), Inches(5.7), Inches(4.3), border_color=SUCCESS_GREEN)
    tb_a8 = s8.shapes.add_textbox(Inches(7.1), Inches(1.65), Inches(5.1), Inches(3.8))
    tf_a8 = tb_a8.text_frame
    tf_a8.word_wrap = True
    p = tf_a8.paragraphs[0]
    p.text = "After: 스키마 단순화 + Python 검증 ➔ PASS"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    a_points_v3 = [
        ("출력 스키마 단순화", "Gemini에는 단순화된 JSON 스키마만 전송 (토큰 선택 지시)"),
        ("토큰 카탈로그 주입", "서버가 수집한 사실(K0001~)을 프롬프트로 전달, 모델은 토큰만 선택"),
        ("Python 엄격 검증 파이프라인", "Gemini 출력 ➔ Python DiagnosisDraft 변환 ➔ Pydantic ➔ 결정적 Evidence 검증")
    ]
    for title, desc in a_points_v3:
        p = tf_a8.add_paragraph()
        p.text = "✔ " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_CYAN
        p.space_before = Pt(10)
        p_d = tf_a8.add_paragraph()
        p_d.text = "   " + desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT

    add_card(s8, Inches(0.8), Inches(5.95), Inches(11.733), Inches(0.8), bg_color=CARD_BG_ALT, border_color=PRIMARY_CYAN)
    tb_sum8 = s8.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(11.3), Inches(0.7))
    tf_s8 = tb_sum8.text_frame
    p_s8 = tf_s8.paragraphs[0]
    p_s8.text = "💡 핵심 설계: 모델을 맹신하는 구조가 아니라, 모델의 생성 결과를 서버가 엄격하게 검증하는 구조로 신뢰성을 확보했습니다."
    p_s8.font.name = FONT_FAMILY
    p_s8.font.size = Pt(11)
    p_s8.font.color.rgb = TEXT_WHITE
    p_s8.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 9: Real Server E2E: T01 vs T03 (v3 Verified)
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "실제 Gemini API E2E 검증 결과: T01 성공과 T03 현황")

    add_card(s9, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.3), border_color=SUCCESS_GREEN)
    tb_t01 = s9.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_01 = tb_t01.text_frame
    tf_01.word_wrap = True
    p = tf_01.paragraphs[0]
    p.text = "T01 정상 시나리오: 실서버 E2E 성공 (PASS)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    t01_bullets_v3 = [
        ("검사 식별자", "INSP-001396 (정상 병 검사)"),
        ("FastAPI 응답", "HTTP 200 OK (analysis_status: COMPLETED)"),
        ("Gemini 외부 요청", "2회 모두 HTTP 200 (전체 시간 약 13.578초)"),
        ("토큰 사용량", "총 4,706 토큰 (입력 3,115, 출력 125, Thinking 1,466)"),
        ("도구 실행", "get_inspection_result 실행 (2.23ms, E0001)"),
        ("근거 검증", "관측 사실 6개, 최종 근거 검증 통과")
    ]
    for k, v in t01_bullets_v3:
        p = tf_01.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(8)

    add_card(s9, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.3), border_color=DANGER_CORAL)
    tb_t03 = s9.shapes.add_textbox(Inches(7.1), Inches(1.65), Inches(5.1), Inches(4.8))
    tf_03 = tb_t03.text_frame
    tf_03.word_wrap = True
    p = tf_03.paragraphs[0]
    p.text = "T03 조도 저하 시나리오: 실서버 재검증 (FAIL)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = DANGER_CORAL

    t03_bullets_v3 = [
        ("검사 식별자", "INSP-001398 (조도 저하 주입 검사)"),
        ("도구 1, 2 호출", "요청 1(결과) 및 요청 2(품질) 실서버 HTTP 200 성공"),
        ("조기 종료 로직", "경고 조건부 조사 종료 성공 (불필요한 3차 도구 호출 방지)"),
        ("최종 단순 JSON 생성", "요청 3: HTTP 504 / DEADLINE_EXCEEDED (14.16초)"),
        ("FastAPI 응답", "HTTP 502 (analysis_status: FAILED, 최종 검증 NOT TESTED)"),
        ("504 원인", "Google 서버 내부 원인 UNKNOWN (추가 검증 필요)")
    ]
    for k, v in t03_bullets_v3:
        p = tf_03.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(8)

    p_keep = tf_03.add_paragraph()
    p_keep.text = "※ 판정 보존: Agent 실패에도 기존 판정(REINSPECT)은 안전 유지됨"
    p_keep.font.name = FONT_FAMILY
    p_keep.font.size = Pt(9.5)
    p_keep.font.color.rgb = WARNING_AMBER
    p_keep.space_before = Pt(10)

    # ==========================================
    # SLIDE 10: Real T01 Demo (v3 Correct Schema)
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "시스템 시연: 검사 생성부터 Gemini AI 진단까지")

    add_card(s10, Inches(0.8), Inches(1.4), Inches(4.5), Inches(5.3))
    tb_d10 = s10.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(3.9), Inches(4.8))
    tf_d10 = tb_d10.text_frame
    tf_d10.word_wrap = True
    p = tf_d10.paragraphs[0]
    p.text = "시연 구성 (재현 가능)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_CYAN

    d_steps_v3 = [
        ("Step 1. 시스템 헬스체크", "GET /health ➔ 모델 적재 확인"),
        ("Step 2. 검사 실행", "POST /inspections ➔ IMG-BOTTLE-0084 투입 (점수 31.33 ➔ OK)"),
        ("Step 3. AI 진단 결과 확인", "INSP-001396 실서버 성공 응답 확인 (근거 매핑 및 요약)")
    ]
    for stitle, sdesc in d_steps_v3:
        p = tf_d10.add_paragraph()
        p.text = stitle
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(14)
        p_d = tf_d10.add_paragraph()
        p_d.text = "  " + sdesc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_MUTED

    p_warn10 = tf_d10.add_paragraph()
    p_warn10.text = "※ 발표 원칙: 이미 저장된 성공 응답과 읽기 전용 API를 중심으로 시연 (유료 API 실시간 호출 지양)"
    p_warn10.font.name = FONT_FAMILY
    p_warn10.font.size = Pt(9.5)
    p_warn10.font.color.rgb = WARNING_AMBER
    p_warn10.space_before = Pt(20)

    add_card(s10, Inches(5.6), Inches(1.4), Inches(6.933), Inches(5.3), border_color=SUCCESS_GREEN)
    tb_j10 = s10.shapes.add_textbox(Inches(5.9), Inches(1.65), Inches(6.3), Inches(4.8))
    tf_j10 = tb_j10.text_frame
    tf_j10.word_wrap = True
    p = tf_j10.paragraphs[0]
    p.text = "실제 T01 실서버 성공 진단 보고서 (INSP-001396)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    json_highlights_v3 = [
        ("overall_summary", "\"기록된 모델 판정과 종합 판정은 OK입니다. 실제 무결점 제품을 보증하지 않습니다.\"", TEXT_LIGHT),
        ("observed_facts (6개)", "F0001~F0006 모두 evidence_id: E0001 및 JSON Pointer(/data/status 등) 매핑", SUCCESS_GREEN),
        ("판정 보존", "inspection_result: OK, final_result: OK (기존 검사 판정 보존)", PRIMARY_CYAN),
        ("recommended_checks (1개)", "check_code: CONTINUE_MONITORING (evidence_refs: F0001, F0002, F0003)", TEXT_LIGHT),
        ("possible_causes", "빈 배열 [] (확인되지 않은 원인 미지목)", TEXT_MUTED),
        ("limitations (4개)", "서버 안전 정책에 따른 한계 고지 4건 포함", WARNING_AMBER)
    ]
    for k, v, c in json_highlights_v3:
        p = tf_j10.add_paragraph()
        p.text = f"• {k}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = c
        p.space_before = Pt(6)
        p_sub = tf_j10.add_paragraph()
        p_sub.text = f"   {v}"
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(10)
        p_sub.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 11: Implementation Status vs Roadmap (v3 718 tests)
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "구현 현황 및 향후 계획 (투명한 성과 구분)")

    col_width = Inches(3.7)
    cols_data_v3 = [
        ("1. 구현 및 실서버 검증",
         SUCCESS_GREEN,
         ["FastAPI 백엔드 & 고유 ID 발급",
          "PatchCore CPU 추론 엔진 (42.5ms)",
          "규칙 기반 종합 판정 (REINSPECT 분리)",
          "Gemini 4대 조회 도구 바인딩",
          "결정적 근거 검증 (Evidence/Pointer)",
          "단순 JSON 스키마 호환 구조",
          "T01 실서버 Gemini E2E 성공 (PASS)"]),

        ("2. Mock 검증 완료 / 실서버 FAIL",
         PRIMARY_CYAN,
         ["전체 pytest 718개 PASS (경고 2건)",
          "T03 경고 조기 종료 로직 통합",
          "T05 로드 실패 / T06 미탐 Mock 검증",
          "판정 불변 가드레일 오프라인 검증",
          "T03 실서버 재검증: FAIL (504 타임아웃)",
          "(T05/T06 실제 실서버는 미검증)"]),

        ("3. 향후 개발 예정",
         WARNING_AMBER,
         ["React 웹 관제 대시보드 (Day 6)",
          "컨베이어 주행 동영상 Live View",
          "실시간 프레임 Grab & 마우스 ROI",
          "ROI 크롭 이미지 PatchCore 검사",
          "Google Cloud Run 컨테이너 배포 (Day 7)",
          "T03/T05/T06 실서버 E2E 재검증"])
    ]
    cx = 0.8
    for ctitle, ccol, citems in cols_data_v3:
        add_card(s11, Inches(cx), Inches(1.4), col_width, Inches(5.3), border_color=ccol)
        tb_col = s11.shapes.add_textbox(Inches(cx + 0.2), Inches(1.65), col_width - Inches(0.4), Inches(4.8))
        tf_col = tb_col.text_frame
        tf_col.word_wrap = True
        p = tf_col.paragraphs[0]
        p.text = ctitle
        p.font.name = FONT_FAMILY
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ccol

        for item in citems:
            p_i = tf_col.add_paragraph()
            p_i.text = "• " + item
            p_i.font.name = FONT_FAMILY
            p_i.font.size = Pt(10.5)
            p_i.font.color.rgb = TEXT_LIGHT
            p_i.space_before = Pt(8)

        cx += 4.0

    # ==========================================
    # SLIDE 12: Limitations & Future Work (v3 Verified)
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "한계 및 향후 개선 계획: 솔직한 엔지니어링 고찰")

    limits_v3 = [
        ("1. PatchCore 미탐 한계 (Recall 50.8%)",
         DANGER_CORAL,
         "현재 ResNet-18 모델은 83장 테스트셋에서 미세 파손/오염 31건을 미검출함 (T06 합격 사례 존재).",
         "개선 계획: 더 큰 백본(WideResNet-50 등)이나 다해상도 패치 결합을 개선 후보로 검토 (성능 보장 아님)"),

        ("2. Gemini HTTP 504 및 보고서 생성 안정성",
         WARNING_AMBER,
         "T03 실서버 재검증에서 최종 JSON 생성 시 14초 후 504 타임아웃 발생 (Google 서버 내부 원인 UNKNOWN).",
         "개선 계획: 프롬프트 압축, timeout 및 재시도 정책 개선은 실제 검증과 운영 안전성 확인이 필요한 계획"),

        ("3. 외부 LLM 호출 지연 (약 13.58초)",
         PRIMARY_CYAN,
         "Gemini API 호출에 약 13초 소요되어 초고속 인라인 실시간 차단에는 부적합.",
         "개선 계획: 비동기 작업 큐(Celery/Redis 등)는 검토 가능한 미래 확장안이며 현재 미구현")
    ]
    ly = 1.4
    for ltitle, lcol, lprob, lsol in limits_v3:
        add_card(s12, Inches(0.8), Inches(ly), Inches(11.733), Inches(1.6), border_color=lcol)
        tb_l = s12.shapes.add_textbox(Inches(1.1), Inches(ly + 0.12), Inches(11.1), Inches(1.4))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True
        p = tf_l.paragraphs[0]
        p.text = ltitle
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = lcol

        p_p = tf_l.add_paragraph()
        p_p.text = "• 한계: " + lprob
        p_p.font.name = FONT_FAMILY
        p_p.font.size = Pt(10.5)
        p_p.font.color.rgb = TEXT_LIGHT
        p_p.space_before = Pt(4)

        p_s = tf_l.add_paragraph()
        p_s.text = "✔ 검토: " + lsol
        p_s.font.name = FONT_FAMILY
        p_s.font.size = Pt(10.5)
        p_s.font.color.rgb = SUCCESS_GREEN
        p_s.space_before = Pt(2)

        ly += 1.8

    # ==========================================
    # SLIDE 13: Conclusion (v3 Title & Message)
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13)

    add_card(s13, Inches(0.8), Inches(1.1), Inches(11.733), Inches(5.4), bg_color=CARD_BG_ALT, border_color=PRIMARY_CYAN)

    tb_end = s13.shapes.add_textbox(Inches(1.3), Inches(1.5), Inches(10.7), Inches(4.5))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True

    p = tf_end.paragraphs[0]
    p.text = "VisionInsight — 근거 기반 산업 비전 AI Agent 설계"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_CYAN

    conclusions_v3 = [
        ("PatchCore", "이상 점수 산출 및 모델 판정 소견(result) 제공"),
        ("FastAPI", "서버 규칙 기반 최종 검사 판정(final_result)과 감사 로그 기록"),
        ("Gemini", "도구 조회를 통한 근거 기반 진단과 점검 권고 제안"),
        ("Python", "엄격한 근거 검증과 판정 보호 정책으로 시스템 신뢰성 확보")
    ]
    for ctitle, cdesc in conclusions_v3:
        p_c = tf_end.add_paragraph()
        p_c.text = f"✔ {ctitle}: {cdesc}"
        p_c.font.name = FONT_FAMILY
        p_c.font.size = Pt(12)
        p_c.font.bold = True
        p_c.font.color.rgb = TEXT_LIGHT
        p_c.space_before = Pt(12)

    p_msg = tf_end.add_paragraph()
    p_msg.text = "\"AI가 검사 판정을 대신하는 것이 아니라, 검증된 검사 정보를 운영자가 이해하고 활용하도록 돕는 시스템\""
    p_msg.font.name = FONT_FAMILY
    p_msg.font.size = Pt(13)
    p_msg.font.bold = True
    p_msg.font.color.rgb = PRIMARY_CYAN
    p_msg.space_before = Pt(16)
    p_msg.alignment = PP_ALIGN.CENTER

    p_qa = tf_end.add_paragraph()
    p_qa.text = "경청해 주셔서 감사합니다.  |  Q&A (질의응답)"
    p_qa.font.name = FONT_FAMILY
    p_qa.font.size = Pt(18)
    p_qa.font.bold = True
    p_qa.font.color.rgb = SUCCESS_GREEN
    p_qa.space_before = Pt(16)
    p_qa.alignment = PP_ALIGN.CENTER

    output_path = os.path.join("presentation", "VisionInsight_Presentation_v3.pptx")
    prs.save(output_path)
    print(f"Presentation v3 saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation_v3()
