"""VisionInsight 이미지 검사 검출력 보강 슬라이드 2장 생성 스크립트
Slide 1: 현재 이미지 검사의 기술적 한계
Slide 2: 이미지 검사 검출력 향상 전략 (A·B·C 전략)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_image_improvement_slides(output_path="presentation/VisionInsight_Image_Inspection_Improvement_Slides.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette (기존 14장 슬라이드와 100% 동일한 토큰)
    BG_COLOR = RGBColor(12, 19, 34)          # #0C1322
    CARD_BG = RGBColor(22, 30, 46)           # #161E2E
    CARD_BG_ALT = RGBColor(16, 23, 38)       # #101726
    CARD_BORDER = RGBColor(42, 54, 73)       # #2A3649
    PRIMARY_CYAN = RGBColor(76, 215, 246)    # #4CD7F6
    SUCCESS_GREEN = RGBColor(52, 211, 153)   # #34D399
    WARNING_AMBER = RGBColor(251, 191, 36)   # #FBBF24
    DANGER_CORAL = RGBColor(248, 113, 113)   # #F87171
    TEXT_WHITE = RGBColor(248, 250, 252)     # #F8FAFC
    TEXT_LIGHT = RGBColor(203, 213, 225)     # #CBD5E1
    TEXT_MUTED = RGBColor(148, 163, 184)     # #94A3B8

    FONT_FAMILY = "Malgun Gothic"

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR
        return bg

    def add_header(slide, title_text, category_text="VISIONINSIGHT  |  IMAGE INSPECTION & QUALITY LIMITATIONS", subtitle_text=None):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.36), Inches(11.7), Inches(0.25))
        tf = cat_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = category_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_CYAN

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.62), Inches(11.7), Inches(0.68))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_top = tf_t.margin_bottom = tf_t.margin_left = tf_t.margin_right = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(20)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

        if subtitle_text:
            p_sub = tf_t.add_paragraph()
            p_sub.text = subtitle_text
            p_sub.font.name = FONT_FAMILY
            p_sub.font.size = Pt(11)
            p_sub.font.color.rgb = TEXT_LIGHT
            p_sub.space_before = Pt(2)

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    # ==========================================
    # SLIDE 1: 현재 이미지 검사의 기술적 한계
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    add_header(s1, "현재 이미지 검사의 기술적 한계",
               category_text="VISIONINSIGHT  |  IMAGE INSPECTION & QUALITY LIMITATIONS",
               subtitle_text="학습 데이터와 실시간 Grab 환경 차이, 단일 평균 밝기 지표, 상대 정규화 Heatmap으로 인한 해석 한계")

    # Left Column: 학습 이미지 vs 실시간 Grab 이미지 비교 (너비 5.7인치, 높이 5.15인치)
    add_card(s1, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.15), border_color=DANGER_CORAL)
    tb_l = s1.shapes.add_textbox(Inches(1.05), Inches(1.6), Inches(5.2), Inches(4.75))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "1. 학습 이미지 ↔ 실시간 LiveView Grab 환경 차이"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = DANGER_CORAL

    diff_items = [
        ("배경 밝기 및 조명 조건의 불일치",
         "정적 MVTec 원본은 균일하고 밝은 배경이나, 동영상 Grab은 컨베이어 배경 명도와 조명 반사광이 상이함"),
        ("ROI 구성 및 제품 위치·크기 변동",
         "정적 영상(900×900 중앙 정렬)과 달리, LiveView Crop은 제품 위치·기울기 및 배경 여백 비율이 가변적임"),
        ("정상 제품의 영상 품질 경고 발생",
         "실제 결함이 없는 정상 제품임에도 환경 차이로 인해 평균 GV 120 전후의 'WARNING_LOW_BRIGHTNESS'가 관측됨")
    ]
    for title, desc in diff_items:
        p = tf_l.add_paragraph()
        p.text = "• " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(8)
        p_d = tf_l.add_paragraph()
        p_d.text = "   " + desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_LIGHT

    # Right Top Column: 평균 밝기(GV) 평가와 서버 판정 영향 (너비 5.733인치, 높이 2.45인치)
    add_card(s1, Inches(6.8), Inches(1.4), Inches(5.733), Inches(2.45), border_color=WARNING_AMBER)
    tb_rt = s1.shapes.add_textbox(Inches(7.05), Inches(1.55), Inches(5.233), Inches(2.15))
    tf_rt = tb_rt.text_frame
    tf_rt.word_wrap = True
    p = tf_rt.paragraphs[0]
    p.text = "2. 평균 밝기(GV) 평가의 구조적 한계"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = WARNING_AMBER

    gv_items = [
        ("전체 ROI 단순 평균의 왜곡",
         "현재 125.0 기준은 ROI 전체 평균치를 측정하므로, 어두운 병 내부와 밝은 배경 비율에 따라 크게 변동"),
        ("서버 최종 REINSPECT 판정 유발 흐름",
         "Grab 이미지 → 평균 GV 평가 (125 미만 시 경고) → PatchCore OK라도 서버가 'REINSPECT' 판정 확정")
    ]
    for title, desc in gv_items:
        p = tf_rt.add_paragraph()
        p.text = "• " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(5)
        p_d = tf_rt.add_paragraph()
        p_d.text = "   " + desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_LIGHT

    # Right Bottom Column: PatchCore 이상지도(Heatmap) 해석 한계 (너비 5.733인치, 높이 2.55인치)
    add_card(s1, Inches(6.8), Inches(4.0), Inches(5.733), Inches(2.55), border_color=PRIMARY_CYAN)
    tb_rb = s1.shapes.add_textbox(Inches(7.05), Inches(4.15), Inches(5.233), Inches(2.25))
    tf_rb = tb_rb.text_frame
    tf_rb.word_wrap = True
    p = tf_rb.paragraphs[0]
    p.text = "3. PatchCore 이상지도(Heatmap) 해석 시 주의점"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_CYAN

    hm_items = [
        ("제품 밖 배경 영역의 이상 반응",
         "모델 특성상 제품 경계 밖 흰 배경 모서리에도 이상 점수가 존재하여 배경 반응 발생"),
        ("이미지별 상대 정규화(Min-Max)의 착시",
         "각 이미지의 최솟값·최댓값 기준으로 색상을 매핑하므로 미세 편차도 시각적으로 붉게 강조됨"),
        ("빨간 영역 ≠ 실제 불량 확정",
         "붉은색 표시는 상대적 이상도 분포일 뿐이며, 단독으로 물리적 결함을 확정하지 않음")
    ]
    for title, desc in hm_items:
        p = tf_rb.add_paragraph()
        p.text = "• " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(4)
        p_d = tf_rb.add_paragraph()
        p_d.text = "   " + desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_LIGHT

    # Bottom Banner (검증 사실 관계 명시)
    add_card(s1, Inches(0.8), Inches(6.7), Inches(11.733), Inches(0.55), bg_color=CARD_BG_ALT, border_color=CARD_BORDER)
    tb_fact = s1.shapes.add_textbox(Inches(0.95), Inches(6.75), Inches(11.4), Inches(0.45))
    tf_f = tb_fact.text_frame
    tf_f.word_wrap = True
    p_f = tf_f.paragraphs[0]
    p_f.text = "💡 실증 사실: 영상 Grab, PatchCore 추론, Heatmap 생성, 서버 판정(OK/NG/REINSPECT)은 실제 성공 검증되었으며, 지표의 타당성과 시각화 해석 기준을 고도화하는 과제입니다."
    p_f.font.name = FONT_FAMILY
    p_f.font.size = Pt(10)
    p_f.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 2: 이미지 검사 검출력 향상 전략 (A·B·C 전략)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "이미지 검사 검출력 향상 전략",
               category_text="VISIONINSIGHT  |  ACCURACY ENHANCEMENT STRATEGY",
               subtitle_text="단순 임계값 조정이 아닌, 환경 정규화(A) · 데이터 정합성(B) · 복합 보조 검사(C)의 다각적 접근")

    # 3-Column Card Layout (너비 3.7인치씩, 간격 0.3인치)
    # Card A: 배경·밝기 영향 최소화
    add_card(s2, Inches(0.8), Inches(1.4), Inches(3.7), Inches(4.7), border_color=PRIMARY_CYAN)
    tb_a = s2.shapes.add_textbox(Inches(0.95), Inches(1.55), Inches(3.4), Inches(4.4))
    tf_a = tb_a.text_frame
    tf_a.word_wrap = True
    p = tf_a.paragraphs[0]
    p.text = "A. 배경·밝기 영향 최소화"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_CYAN

    p_sub_a = tf_a.add_paragraph()
    p_sub_a.text = "[촬영 품질 분리 및 정규화]"
    p_sub_a.font.name = FONT_FAMILY
    p_sub_a.font.size = Pt(10)
    p_sub_a.font.color.rgb = TEXT_MUTED
    p_sub_a.space_before = Pt(2)

    card_a_items = [
        ("병 전경/배경 분리 품질 지표 실험",
         "ROI 전체 평균에서 탈피하여 병 내부 GV, 배경 GV, 전경/배경 대비를 분리 측정해 품질 평가 타당성 확보"),
        ("ROI 및 조명 조건 표준화",
         "병 중심 정렬 알고리즘을 도입하고, 컨베이어 조명 반사광 및 Crop 주변 여백을 일정한 비율로 제어"),
        ("실제 Grab 표본 기반 기준 재검토",
         "실제 Grab 정상 샘플의 지표 분포를 통계적으로 분석하여, 현재 운영 가이드(125.0)의 오경보·미경보를 비교 검토")
    ]
    for title, desc in card_a_items:
        p = tf_a.add_paragraph()
        p.text = "• " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)
        p_d = tf_a.add_paragraph()
        p_d.text = "   " + desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_LIGHT

    # Card B: 학습 데이터와 실제 환경 정합성 개선
    add_card(s2, Inches(4.8), Inches(1.4), Inches(3.7), Inches(4.7), border_color=WARNING_AMBER)
    tb_b = s2.shapes.add_textbox(Inches(4.95), Inches(1.55), Inches(3.4), Inches(4.4))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "B. 환경 정합성 개선"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = WARNING_AMBER

    p_sub_b = tf_b.add_paragraph()
    p_sub_b.text = "[현장 데이터셋 구축 & 재평가]"
    p_sub_b.font.name = FONT_FAMILY
    p_sub_b.font.size = Pt(10)
    p_sub_b.font.color.rgb = TEXT_MUTED
    p_sub_b.space_before = Pt(2)

    card_b_items = [
        ("Grab 독립 검증 데이터셋 확충",
         "정상 Grab, 결함(파손·이물) Grab, 저품질 Grab을 독립 수집하고, 정답 라벨은 Agent와 엄격히 분리"),
        ("고정 모델 기반 정량 재평가",
         "현재 PatchCore 모델을 고정한 채 혼동행렬, Precision, Recall, 오탐(FP), 미검(FN)을 객관적으로 실측"),
        ("실측 기반 임계값 재보정 검토",
         "임의의 수치 변경을 엄금하고, 독립 평가셋의 ROC/PR 곡선과 유형별 오검·미검 수치에 근거해 임계값(34.1442) 재보정 검토")
    ]
    for title, desc in card_b_items:
        p = tf_b.add_paragraph()
        p.text = "• " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)
        p_d = tf_b.add_paragraph()
        p_d.text = "   " + desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_LIGHT

    # Card C: 보조 검사 기술 결합
    add_card(s2, Inches(8.8), Inches(1.4), Inches(3.7), Inches(4.7), border_color=SUCCESS_GREEN)
    tb_c = s2.shapes.add_textbox(Inches(8.95), Inches(1.55), Inches(3.4), Inches(4.4))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    p = tf_c.paragraphs[0]
    p.text = "C. 보조 검사 기술 결합"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    p_sub_c = tf_c.add_paragraph()
    p_sub_c.text = "[OpenCV 규칙 & 다중 교차 검증]"
    p_sub_c.font.name = FONT_FAMILY
    p_sub_c.font.size = Pt(10)
    p_sub_c.font.color.rgb = TEXT_MUTED
    p_sub_c.space_before = Pt(2)

    card_c_items = [
        ("OpenCV 기하학적 형상 검사 검토",
         "병 입구 원형도, 편심, 림(Rim) 경계 파손, 두께 불균일 등 측정 가능한 기하학적 보조 규칙 병렬 실험"),
        ("제품 영역 분리 기반 이상지도 해석",
         "단순 시각적 배경 지우기가 아닌, 병 마스크를 기반으로 제품 내부/림/배경의 영역별 이상도 통계를 산출"),
        ("Gemini Vision Judge 2차 참고 활용",
         "이미 실호출에 성공한 Gemini 시각 분석을 독립적인 2차 참고 소견으로 결합 (서버 판정 자동 대체 금지)")
    ]
    for title, desc in card_c_items:
        p = tf_c.add_paragraph()
        p.text = "• " + title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)
        p_d = tf_c.add_paragraph()
        p_d.text = "   " + desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_LIGHT

    # Bottom Banner: 단계별 검증 파이프라인 및 기대 효과
    add_card(s2, Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.9), bg_color=CARD_BG_ALT, border_color=CARD_BORDER)
    tb_bot = s2.shapes.add_textbox(Inches(0.95), Inches(6.3), Inches(11.4), Inches(0.8))
    tf_bot = tb_bot.text_frame
    tf_bot.word_wrap = True
    p_b0 = tf_bot.paragraphs[0]
    p_b0.text = "단계별 검증 로드맵: P0 (기준선 데이터 수집)  ➔  P1 (A 품질 지표 & B 검증셋 FP/FN 측정)  ➔  P2 (C 보조 규칙 병렬 결합)"
    p_b0.font.name = FONT_FAMILY
    p_b0.font.size = Pt(10.5)
    p_b0.font.bold = True
    p_b0.font.color.rgb = PRIMARY_CYAN

    p_b1 = tf_bot.add_paragraph()
    p_b1.text = "🎯 기대 효과: 환경 차이로 인한 불필요한 재검사(REINSPECT) 오경보 최소화  &  미세 결함 미검(FN) 위험 완화 (※ 과장된 수치 대신 실측 비교 원칙 준수)"
    p_b1.font.name = FONT_FAMILY
    p_b1.font.size = Pt(10)
    p_b1.font.color.rgb = TEXT_LIGHT
    p_b1.space_before = Pt(3)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_image_improvement_slides()
