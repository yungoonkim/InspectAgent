"""VisionInsight 이미지 검사 검출력 및 Grab 화질 개선 슬라이드 2장 생성 스크립트 (ver2)
Slide 1: 현재 이미지 검사의 기술적 한계 (환경 차이, 평균 GV, Heatmap 상대 정규화, Grab 화질 흐림)
Slide 2: 이미지 검사 검출력 향상 전략 (A·B·C 전략 및 이번 주 개선 계획)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_image_improvement_slides_v2(output_path="presentation/VisionInsight_Image_Inspection_Improvement_Slides_v2.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette (기존 VisionInsight 디자인 시스템과 100% 동일)
    BG_COLOR = RGBColor(12, 19, 34)          # #0C1322
    CARD_BG = RGBColor(22, 30, 46)           # #161E2E
    CARD_BG_ALT = RGBColor(16, 23, 38)       # #101726
    CARD_BORDER = RGBColor(42, 54, 73)       # #2A3649
    PRIMARY_CYAN = RGBColor(76, 215, 246)    # #4CD7F6
    SUCCESS_GREEN = RGBColor(52, 211, 153)   # #34D399
    WARNING_AMBER = RGBColor(251, 191, 36)   # #FBBF24
    DANGER_CORAL = RGBColor(248, 113, 113)   # #F87171
    TEXT_WHITE = RGBColor(248, 250, 252)     # #F8FAFC
    TEXT_LIGHT = RGBColor(203, 213, 225)     # #CBD5E1
    TEXT_MUTED = RGBColor(148, 163, 184)     # #94A3B8

    FONT_FAMILY = "Malgun Gothic"

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR
        return bg

    def add_header(slide, title_text, category_text="VISIONINSIGHT  |  IMAGE INSPECTION & QUALITY LIMITATIONS", subtitle_text=None):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.34), Inches(11.7), Inches(0.24))
        tf = cat_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = category_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_CYAN

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.58), Inches(11.7), Inches(0.68))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_top = tf_t.margin_bottom = tf_t.margin_left = tf_t.margin_right = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(20)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

        if subtitle_text:
            p_sub = tf_t.add_paragraph()
            p_sub.text = subtitle_text
            p_sub.font.name = FONT_FAMILY
            p_sub.font.size = Pt(11)
            p_sub.font.color.rgb = TEXT_LIGHT
            p_sub.space_before = Pt(2)

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    # ==========================================
    # SLIDE 1: 현재 이미지 검사의 기술적 한계 (4개 영역 2x2 그리드)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    add_header(s1, "현재 이미지 검사의 기술적 한계",
               category_text="VISIONINSIGHT  |  IMAGE INSPECTION & QUALITY LIMITATIONS",
               subtitle_text="학습-실검사 환경 차이 · 평균 밝기 지표 종속 · 상대 정규화 Heatmap · Grab 세부 화질 저하 분석")

    col_w = Inches(5.72)
    card_h = Inches(2.45)
    left_col_x = Inches(0.8)
    right_col_x = Inches(6.81)
    row1_y = Inches(1.35)
    row2_y = Inches(3.95)

    # 1. 학습 이미지 vs LiveView Grab 환경 차이 (좌측 상단)
    add_card(s1, left_col_x, row1_y, col_w, card_h, border_color=WARNING_AMBER)
    tb_1 = s1.shapes.add_textbox(left_col_x + Inches(0.18), row1_y + Inches(0.12), col_w - Inches(0.36), card_h - Inches(0.2))
    tf_1 = tb_1.text_frame
    tf_1.word_wrap = True
    p1 = tf_1.paragraphs[0]
    p1.text = "1. 학습 이미지 ↔ 실시간 LiveView Grab 환경 차이"
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(13.5)
    p1.font.bold = True
    p1.font.color.rgb = WARNING_AMBER

    items_1 = [
        ("배경 조명 및 촬영 조건 불일치", "정적 MVTec 원본은 균일한 밝은 단색 배경이나, LiveView Grab은 컨베이어 배경 명도와 조명 반사광이 상이함"),
        ("ROI 구도 및 제품 위치·크기 변동", "정적 영상(900×900 중앙 정렬)과 달리, LiveView Crop은 제품 위치·기울기 및 배경 여백 비율이 가변적임"),
        ("정상 제품의 영상 품질 경고 유발", "실제 결함이 없는 양품임에도 환경 차이로 인해 평균 GV 120 전후의 'WARNING_LOW_BRIGHTNESS' 관측")
    ]
    for t, d in items_1:
        p = tf_1.add_paragraph()
        p.text = "• " + t
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(4)
        pd = tf_1.add_paragraph()
        pd.text = "   " + d
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_LIGHT

    # 2. 평균 밝기(GV) 평가의 구조적 한계 (우측 상단)
    add_card(s1, right_col_x, row1_y, col_w, card_h, border_color=DANGER_CORAL)
    tb_2 = s1.shapes.add_textbox(right_col_x + Inches(0.18), row1_y + Inches(0.12), col_w - Inches(0.36), card_h - Inches(0.2))
    tf_2 = tb_2.text_frame
    tf_2.word_wrap = True
    p2 = tf_2.paragraphs[0]
    p2.text = "2. 평균 밝기(GV) 평가의 구조적 한계"
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(13.5)
    p2.font.bold = True
    p2.font.color.rgb = DANGER_CORAL

    items_2 = [
        ("전체 ROI 단순 평균치의 왜곡", "현재 125.0 기준은 ROI 전체 평균치를 측정하므로, 어두운 병 내부와 밝은 배경 비율에 따라 크게 변동"),
        ("서버 최종 REINSPECT 판정 유발 흐름", "Grab 이미지 → 평균 GV 평가 (125 미만 시 경고) → PatchCore OK라도 서버가 'REINSPECT' 판정 확정"),
        ("품질 평가 지표 분리 필요성", "조명 결함과 제품 형상을 구분하기 위해 병 전경과 배경을 분리한 독립 평가 체계가 필수적임")
    ]
    for t, d in items_2:
        p = tf_2.add_paragraph()
        p.text = "• " + t
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(4)
        pd = tf_2.add_paragraph()
        pd.text = "   " + d
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_LIGHT

    # 3. PatchCore 이상지도(Heatmap) 해석 시 주의점 (좌측 하단)
    add_card(s1, left_col_x, row2_y, col_w, card_h, border_color=PRIMARY_CYAN)
    tb_3 = s1.shapes.add_textbox(left_col_x + Inches(0.18), row2_y + Inches(0.12), col_w - Inches(0.36), card_h - Inches(0.2))
    tf_3 = tb_3.text_frame
    tf_3.word_wrap = True
    p3 = tf_3.paragraphs[0]
    p3.text = "3. PatchCore 이상지도(Heatmap) 해석 주의점"
    p3.font.name = FONT_FAMILY
    p3.font.size = Pt(13.5)
    p3.font.bold = True
    p3.font.color.rgb = PRIMARY_CYAN

    items_3 = [
        ("제품 밖 배경 영역의 이상 반응", "모델 특성상 제품 경계 밖 흰 배경 모서리 영역에도 이상 점수가 수학적으로 존재하여 배경 반응 발생"),
        ("이미지별 상대 정규화(Min-Max) 착시", "각 이미지의 최솟값·최댓값 기준으로 색상을 매핑하므로 미세 편차도 시각적으로 붉게 강조됨"),
        ("빨간 영역 ≠ 실제 불량 확정", "붉은색 표시는 상대적 이상도 분포일 뿐이며, 단독으로 물리적 결함을 확정하지 않음 (빨강 ≠ 결함)")
    ]
    for t, d in items_3:
        p = tf_3.add_paragraph()
        p.text = "• " + t
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(4)
        pd = tf_3.add_paragraph()
        pd.text = "   " + d
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_LIGHT

    # 4. LiveView Grab 세부 화질 저하 현상 (우측 하단 - 신규 보강)
    add_card(s1, right_col_x, row2_y, col_w, card_h, border_color=PRIMARY_CYAN)
    tb_4 = s1.shapes.add_textbox(right_col_x + Inches(0.18), row2_y + Inches(0.12), col_w - Inches(0.36), card_h - Inches(0.2))
    tf_4 = tb_4.text_frame
    tf_4.word_wrap = True
    p4 = tf_4.paragraphs[0]
    p4.text = "4. 원본 대비 LiveView Grab 세부 화질 저하 현상"
    p4.font.name = FONT_FAMILY
    p4.font.size = Pt(13.5)
    p4.font.bold = True
    p4.font.color.rgb = PRIMARY_CYAN

    items_4 = [
        ("실제 이미지 비교 관찰 현상", "정적 원본 대비 LiveView Grab Crop에서 병 림(Rim) 윤곽선과 미세 표면 무늬가 흐릿하게 관측됨"),
        ("원인 후보 분리 분석 (미확정)", "화질 저하 경로: 원본 소스 → 영상 합성 → H.264 인코딩 → 디코딩·Grab → Crop 저장 → 뷰어 확대"),
        ("검증 전 단정 금지 원칙", "초점 불량으로 단정하지 않으며, 각 단계별(압축·리사이즈·보간) 정보 손실 기여도를 정량 측정 예정")
    ]
    for t, d in items_4:
        p = tf_4.add_paragraph()
        p.text = "• " + t
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(4)
        pd = tf_4.add_paragraph()
        pd.text = "   " + d
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_LIGHT

    # Bottom Banner (검증 사실 vs 미검증 원인 분리)
    add_card(s1, Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.6), bg_color=CARD_BG_ALT, border_color=CARD_BORDER)
    tb_fact = s1.shapes.add_textbox(Inches(0.95), Inches(6.6), Inches(11.4), Inches(0.5))
    tf_f = tb_fact.text_frame
    tf_f.word_wrap = True
    p_f = tf_f.paragraphs[0]
    p_f.text = "💡 실증 사실과 과제: 영상 Grab, PatchCore 추론, Heatmap 생성, 서버 판정(OK/NG/REINSPECT)은 실제 성공 검증되었습니다. 다만 조도 경고, 상대 정규화 착시, Grab 세부 흐림은 실측 데이터를 통한 단계별 정량 규명 과제로 관리합니다."
    p_f.font.name = FONT_FAMILY
    p_f.font.size = Pt(9.5)
    p_f.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 2: 이미지 검사 검출력 향상 전략 (A·B·C 전략 및 이번 주 개선 계획)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "이미지 검사 검출력 향상 전략",
               category_text="VISIONINSIGHT  |  ACCURACY ENHANCEMENT STRATEGY",
               subtitle_text="품질·화질 개선(A) · 현장 데이터 정합성(B) · 복합 보조 검사(C) 및 이번 주 실행 계획")

    col_w3 = Inches(3.7)
    card_top_y = Inches(1.35)
    card_top_h = Inches(3.85)

    # Card A: 이미지 품질·촬영 환경 개선 (화질 개선 포함)
    add_card(s2, Inches(0.8), card_top_y, col_w3, card_top_h, border_color=PRIMARY_CYAN)
    tb_a = s2.shapes.add_textbox(Inches(0.95), card_top_y + Inches(0.12), Inches(3.4), card_top_h - Inches(0.2))
    tf_a = tb_a.text_frame
    tf_a.word_wrap = True
    pa = tf_a.paragraphs[0]
    pa.text = "A. 이미지 품질·촬영 환경 개선"
    pa.font.name = FONT_FAMILY
    pa.font.size = Pt(14)
    pa.font.bold = True
    pa.font.color.rgb = PRIMARY_CYAN

    pa_sub = tf_a.add_paragraph()
    pa_sub.text = "[전경/배경 분리 & 선명도 개선]"
    pa_sub.font.name = FONT_FAMILY
    pa_sub.font.size = Pt(9.5)
    pa_sub.font.color.rgb = TEXT_MUTED
    pa_sub.space_before = Pt(2)

    card_a_items = [
        ("병 영역/배경 분리 품질 지표", "병 내부 GV, 배경 GV, 대비를 분리 측정하여 불필요한 REINSPECT 오경보 차단"),
        ("ROI 및 조명 조건 표준화", "병 중심 정렬 알고리즘 도입, 반사광 제어 및 Crop 주변 여백 비율 안정화"),
        ("Grab 화질·선명도 개선", "원본·영상·Crop 단계별 해상도 추적 및 병 림 라플라시안 분산 지표 정량 비교"),
        ("미세 결함 표현력 비교", "동일 결함 표본으로 원본 vs Grab의 작은 크랙·칩핑 가시성 손실 여부 확인")
    ]
    for t, d in card_a_items:
        p = tf_a.add_paragraph()
        p.text = "• " + t
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(6)
        pd = tf_a.add_paragraph()
        pd.text = "   " + d
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_LIGHT

    # Card B: 검사 데이터와 모델 평가 개선
    add_card(s2, Inches(4.8), card_top_y, col_w3, card_top_h, border_color=WARNING_AMBER)
    tb_b = s2.shapes.add_textbox(Inches(4.95), card_top_y + Inches(0.12), Inches(3.4), card_top_h - Inches(0.2))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    pb = tf_b.paragraphs[0]
    pb.text = "B. 검사 데이터 & 평가 개선"
    pb.font.name = FONT_FAMILY
    pb.font.size = Pt(14)
    pb.font.bold = True
    pb.font.color.rgb = WARNING_AMBER

    pb_sub = tf_b.add_paragraph()
    pb_sub.text = "[현장 데이터셋 구축 & 재평가]"
    pb_sub.font.name = FONT_FAMILY
    pb_sub.font.size = Pt(9.5)
    pb_sub.font.color.rgb = TEXT_MUTED
    pb_sub.space_before = Pt(2)

    card_b_items = [
        ("Grab 독립 검증 데이터셋 확충", "정상·불량(파손/이물)·저품질 Grab 독립 수집 및 정답 라벨 엄격 분리"),
        ("영상 환경-학습 데이터 차이 분석", "컨베이어 LiveView 영상과 기존 MVTec 학습셋의 조명·배경 도메인 편차 분석"),
        ("고정 모델 기반 정량 재평가", "PatchCore 모델을 고정한 상태에서 Precision, Recall, 오탐(FP), 미검(FN) 실측"),
        ("실측 기반 임계값 재보정 검토", "임의 수치 변경 엄금, 독립 평가셋의 실측 ROC/PR 곡선 근거 34.1442 재보정 검토")
    ]
    for t, d in card_b_items:
        p = tf_b.add_paragraph()
        p.text = "• " + t
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(6)
        pd = tf_b.add_paragraph()
        pd.text = "   " + d
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_LIGHT

    # Card C: 복합 검사 기술 보강
    add_card(s2, Inches(8.8), card_top_y, col_w3, card_top_h, border_color=SUCCESS_GREEN)
    tb_c = s2.shapes.add_textbox(Inches(8.95), card_top_y + Inches(0.12), Inches(3.4), card_top_h - Inches(0.2))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    pc = tf_c.paragraphs[0]
    pc.text = "C. 복합 검사 기술 보강"
    pc.font.name = FONT_FAMILY
    pc.font.size = Pt(14)
    pc.font.bold = True
    pc.font.color.rgb = SUCCESS_GREEN

    pc_sub = tf_c.add_paragraph()
    pc_sub.text = "[Edge 검출 & 다중 교차 검증]"
    pc_sub.font.name = FONT_FAMILY
    pc_sub.font.size = Pt(9.5)
    pc_sub.font.color.rgb = TEXT_MUTED
    pc_sub.space_before = Pt(2)

    card_c_items = [
        ("병 외곽 Edge 및 림 검출", "병 외곽 윤곽선 및 내부 림(Rim) 경계를 검출하여 구조적 형상 특징 추출"),
        ("병 영역 검출 정확도 검증", "병 마스킹 시각화 오버레이 정확도 검증 및 영역별 Heatmap 통계 분리 산출"),
        ("OpenCV 기하학적 보조 검사", "병 입구 원형도, 편심, 림 파손, 테두리 불균일 보조 규칙 병렬 실험"),
        ("Gemini Vision Judge 2차 참고", "실호출 성공한 Gemini 시각 분석을 독립 2차 참고 소견으로 활용 (서버 판정 대체 불가)")
    ]
    for t, d in card_c_items:
        p = tf_c.add_paragraph()
        p.text = "• " + t
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(6)
        pd = tf_c.add_paragraph()
        pd.text = "   " + d
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_LIGHT

    # Bottom Area: 이번 주 개선 계획 (4대 실행 과제 칩 형태)
    bot_y = Inches(5.35)
    bot_h = Inches(1.8)
    add_card(s2, Inches(0.8), bot_y, Inches(11.733), bot_h, bg_color=CARD_BG_ALT, border_color=PRIMARY_CYAN)

    tb_plan_title = s2.shapes.add_textbox(Inches(1.0), bot_y + Inches(0.08), Inches(11.3), Inches(0.3))
    tf_pt = tb_plan_title.text_frame
    tf_pt.word_wrap = True
    ppt = tf_pt.paragraphs[0]
    ppt.text = "🎯 이번 주 핵심 실행 계획 (Weekly Action Plan — 현황 검증 및 오프라인 실험)"
    ppt.font.name = FONT_FAMILY
    ppt.font.size = Pt(11)
    ppt.font.bold = True
    ppt.font.color.rgb = PRIMARY_CYAN

    # 4개 과제 칩 (가로 4열 배치)
    plans = [
        ("1. 병 테두리 검출·오버레이 검증", "Grab 이미지 병 외곽·림 Edge 검출 및 시각화 정확도 측정"),
        ("2. 단계별 선명도 비교·개선", "원본→영상→Crop 단계별 라플라시안 분산 측정 및 손실 분석"),
        ("3. 전경/배경 품질 평가 실험", "병 영역과 배경을 분리한 GV 지표의 오프라인 비교 평가"),
        ("4. 검증 데이터 구축 및 재평가", "정상·불량 Grab 표본 확보 및 PatchCore의 FP/FN 수치 실측")
    ]
    chip_w = Inches(2.7)
    chip_gap = Inches(0.17)
    chip_x = Inches(1.0)
    chip_y = bot_y + Inches(0.42)
    chip_h = Inches(0.85)

    for p_title, p_desc in plans:
        add_card(s2, chip_x, chip_y, chip_w, chip_h, bg_color=CARD_BG, border_color=CARD_BORDER)
        tb_cp = s2.shapes.add_textbox(chip_x + Inches(0.08), chip_y + Inches(0.06), chip_w - Inches(0.16), chip_h - Inches(0.12))
        tf_cp = tb_cp.text_frame
        tf_cp.word_wrap = True
        pcp0 = tf_cp.paragraphs[0]
        pcp0.text = p_title
        pcp0.font.name = FONT_FAMILY
        pcp0.font.size = Pt(9.5)
        pcp0.font.bold = True
        pcp0.font.color.rgb = TEXT_WHITE

        pcp1 = tf_cp.add_paragraph()
        pcp1.text = p_desc
        pcp1.font.name = FONT_FAMILY
        pcp1.font.size = Pt(8.5)
        pcp1.font.color.rgb = TEXT_LIGHT
        pcp1.space_before = Pt(2)

        chip_x += chip_w + chip_gap

    # Footer note inside bottom card
    tb_fn = s2.shapes.add_textbox(Inches(1.0), bot_y + Inches(1.35), Inches(11.3), Inches(0.35))
    tf_fn = tb_fn.text_frame
    pfn = tf_fn.paragraphs[0]
    pfn.text = "※ 위 항목은 이번 주 진행 목표이며 구현 완료나 검출률 향상이 확정된 결과가 아닙니다. 기존 코드·임계값은 근거 없이 변경하지 않습니다."
    pfn.font.name = FONT_FAMILY
    pfn.font.size = Pt(9)
    pfn.font.color.rgb = TEXT_MUTED

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_image_improvement_slides_v2()
