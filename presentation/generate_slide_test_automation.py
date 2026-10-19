"""VisionInsight 테스트 자동화 및 오류 개선 슬라이드 2장 생성 스크립트
Slide 1: Playwright E2E — 자동화 테스트로 실제 오류 발견 및 해결 (120→121→119 오류 및 120→121→120 개선 비교)
Slide 2: 자동화 테스트 현황 및 데이터 보호 기반 검증 계획 (3 PASS / 2 BLOCKED 및 격리 환경 후속 계획)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_test_automation_slides(output_path="presentation/VisionInsight_Test_Automation_Improvement_Slides.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette (VisionInsight 다크 네이비 테마와 100% 동일)
    BG_COLOR = RGBColor(12, 19, 34)          # #0C1322
    CARD_BG = RGBColor(22, 30, 46)           # #161E2E
    CARD_BG_ALT = RGBColor(16, 23, 38)       # #101726
    CARD_BORDER = RGBColor(42, 54, 73)       # #2A3649
    PRIMARY_CYAN = RGBColor(76, 215, 246)    # #4CD7F6
    SUCCESS_GREEN = RGBColor(52, 211, 153)   # #34D399
    WARNING_AMBER = RGBColor(251, 191, 36)   # #FBBF24
    DANGER_CORAL = RGBColor(248, 113, 113)   # #F87171
    TEXT_WHITE = RGBColor(248, 250, 252)     # #F8FAFC
    TEXT_LIGHT = RGBColor(203, 213, 225)     # #CBD5E1
    TEXT_MUTED = RGBColor(148, 163, 184)     # #94A3B8

    FONT_FAMILY = "Malgun Gothic"

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR
        return bg

    def add_header(slide, title_text, category_text="VISIONINSIGHT  |  E2E TEST AUTOMATION & QUALITY ENGINEERING", subtitle_text=None):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.34), Inches(11.7), Inches(0.24))
        tf = cat_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = category_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_CYAN

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.58), Inches(11.7), Inches(0.68))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_top = tf_t.margin_bottom = tf_t.margin_left = tf_t.margin_right = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(20)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

        if subtitle_text:
            p_sub = tf_t.add_paragraph()
            p_sub.text = subtitle_text
            p_sub.font.name = FONT_FAMILY
            p_sub.font.size = Pt(11)
            p_sub.font.color.rgb = TEXT_LIGHT
            p_sub.space_before = Pt(2)

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    # ==========================================
    # SLIDE 1: 자동화 테스트로 실제 오류 발견 및 해결
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    add_header(s1, "Playwright E2E — 자동화 테스트로 실제 오류 발견 및 해결",
               category_text="VISIONINSIGHT  |  E2E TEST AUTOMATION & QUALITY ENGINEERING",
               subtitle_text="Gemini 시나리오 15개 기반 · Playwright 읽기 전용 5대 흐름 구현 및 Chrome 1프레임 왕복 오차 교정")

    # 발표자 노트 (약 30초 대본)
    s1.notes_slide.notes_text_frame.text = (
        "기능 구현 후 수동 시연뿐 아니라 Playwright를 통한 실제 Chrome 자동화 검증을 도입했습니다. "
        "그 과정에서 120프레임에서 한 프레임 앞으로 이동했다가 되돌아오면 119프레임이 표시되는 오류를 포착했습니다. "
        "시간 계산을 정수 프레임 중심으로 수정하고 탐색 중 Grab도 제한했습니다. "
        "같은 E2E 테스트와 단위 테스트 8건을 통해 수정 효과를 입증했습니다."
    )

    # 상단 요약 배너 (4대 핵심 지표)
    add_card(s1, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.82), bg_color=CARD_BG_ALT, border_color=PRIMARY_CYAN)
    tb_top1 = s1.shapes.add_textbox(Inches(0.95), Inches(1.42), Inches(11.4), Inches(0.68))
    tf_top1 = tb_top1.text_frame
    tf_top1.word_wrap = True
    p_top1 = tf_top1.paragraphs[0]
    p_top1.text = "E2E 테스트 시나리오 15개 설계 (Gemini) ➔ 읽기 전용 5대 흐름 자동화 구현 (Playwright) ➔ 실제 Chrome 회귀 검증"
    p_top1.font.name = FONT_FAMILY
    p_top1.font.size = Pt(12)
    p_top1.font.bold = True
    p_top1.font.color.rgb = PRIMARY_CYAN

    p_top1_sub = tf_top1.add_paragraph()
    p_top1_sub.text = "• 대상 영상: VID-000003 (16초·480프레임)  |  • 환경: Windows 11 Chrome 로컬 자동 실행  |  • 방식: 무조작·무변조 원칙 준수"
    p_top1_sub.font.name = FONT_FAMILY
    p_top1_sub.font.size = Pt(9.5)
    p_top1_sub.font.color.rgb = TEXT_LIGHT
    p_top1_sub.space_before = Pt(2)

    col_w1 = Inches(5.72)
    card_h1 = Inches(4.15)
    left_x1 = Inches(0.8)
    right_x1 = Inches(6.81)
    card_y1 = Inches(2.28)

    # 1. 오류 발견 (좌측 카드 - DANGER_CORAL)
    add_card(s1, left_x1, card_y1, col_w1, card_h1, border_color=DANGER_CORAL)
    tb_l = s1.shapes.add_textbox(left_x1 + Inches(0.18), card_y1 + Inches(0.12), col_w1 - Inches(0.36), card_h1 - Inches(0.24))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    pl_title = tf_l.paragraphs[0]
    pl_title.text = "1. 실제 오류 발견: 120 → 121 → 119 프레임 오차"
    pl_title.font.name = FONT_FAMILY
    pl_title.font.size = Pt(13)
    pl_title.font.bold = True
    pl_title.font.color.rgb = DANGER_CORAL

    # 수정 전 시각적 다이어그램 박스
    dia_box_l = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x1 + Inches(0.2), card_y1 + Inches(0.48), col_w1 - Inches(0.4), Inches(0.85))
    dia_box_l.fill.solid()
    dia_box_l.fill.fore_color.rgb = CARD_BG_ALT
    dia_box_l.line.color.rgb = DANGER_CORAL
    dia_box_l.line.width = Pt(1)

    tf_dia_l = dia_box_l.text_frame
    tf_dia_l.word_wrap = True
    p_dl = tf_dia_l.paragraphs[0]
    p_dl.text = "[수정 전]   Frame 120   ──(+1 Frame)──▶   Frame 121   ──(-1 Frame)──▶   Frame 119  ❌"
    p_dl.font.name = FONT_FAMILY
    p_dl.font.size = Pt(10.5)
    p_dl.font.bold = True
    p_dl.font.color.rgb = DANGER_CORAL
    p_dl.alignment = PP_ALIGN.CENTER

    p_dl_sub = tf_dia_l.add_paragraph()
    p_dl_sub.text = "* 기대값(120)과 달리 1프레임 뒤처지는 오프바이원(off-by-one) 오차 발생"
    p_dl_sub.font.name = FONT_FAMILY
    p_dl_sub.font.size = Pt(8.5)
    p_dl_sub.font.color.rgb = TEXT_MUTED
    p_dl_sub.alignment = PP_ALIGN.CENTER
    p_dl_sub.space_before = Pt(2)

    # 오류 상세 불렛
    items_l = [
        ("발견 방식", "Playwright TC-01 자동화 스위트 실제 Chrome 실행 중 포착"),
        ("재현 과정", "VID-000003의 120프레임에서 '+1 Frame' 클릭 후 '-1 Frame' 클릭"),
        ("원인 분석", "Chrome 시간 정밀도 부동소수점 오차(3.999999초)와 초 단위 증감, 프레임 인덱스 내림(Math.floor) 계산의 복합 충돌"),
        ("잠재적 위험", "화면 표시 프레임과 Grab 대상 프레임이 1프레임 어긋나는 정합성 결함 유발")
    ]
    tb_l_text = s1.shapes.add_textbox(left_x1 + Inches(0.18), card_y1 + Inches(1.42), col_w1 - Inches(0.36), card_h1 - Inches(1.5))
    tf_lt = tb_l_text.text_frame
    tf_lt.word_wrap = True
    for i, (k, v) in enumerate(items_l):
        p = tf_lt.paragraphs[0] if i == 0 else tf_lt.add_paragraph()
        p.text = "• " + k + ": " + v
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(4) if i > 0 else Pt(0)

    # 2. 개선 조치 및 재검증 (우측 카드 - SUCCESS_GREEN)
    add_card(s1, right_x1, card_y1, col_w1, card_h1, border_color=SUCCESS_GREEN)
    tb_r = s1.shapes.add_textbox(right_x1 + Inches(0.18), card_y1 + Inches(0.12), col_w1 - Inches(0.36), card_h1 - Inches(0.24))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    pr_title = tf_r.paragraphs[0]
    pr_title.text = "2. 개선 조치: 정수 프레임 통일 & 회귀 검증 성공"
    pr_title.font.name = FONT_FAMILY
    pr_title.font.size = Pt(13)
    pr_title.font.bold = True
    pr_title.font.color.rgb = SUCCESS_GREEN

    # 수정 후 시각적 다이어그램 박스
    dia_box_r = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x1 + Inches(0.2), card_y1 + Inches(0.48), col_w1 - Inches(0.4), Inches(0.85))
    dia_box_r.fill.solid()
    dia_box_r.fill.fore_color.rgb = CARD_BG_ALT
    dia_box_r.line.color.rgb = SUCCESS_GREEN
    dia_box_r.line.width = Pt(1)

    tf_dia_r = dia_box_r.text_frame
    tf_dia_r.word_wrap = True
    p_dr = tf_dia_r.paragraphs[0]
    p_dr.text = "[수정 후]   Frame 120   ──(+1 Frame)──▶   Frame 121   ──(-1 Frame)──▶   Frame 120  ✔"
    p_dr.font.name = FONT_FAMILY
    p_dr.font.size = Pt(10.5)
    p_dr.font.bold = True
    p_dr.font.color.rgb = SUCCESS_GREEN
    p_dr.alignment = PP_ALIGN.CENTER

    p_dr_sub = tf_dia_r.add_paragraph()
    p_dr_sub.text = "* 정수 프레임 인덱스 통일 및 구간 중앙값 Seek로 정확히 120 복귀"
    p_dr_sub.font.name = FONT_FAMILY
    p_dr_sub.font.size = Pt(8.5)
    p_dr_sub.font.color.rgb = TEXT_MUTED
    p_dr_sub.alignment = PP_ALIGN.CENTER
    p_dr_sub.space_before = Pt(2)

    # 수정 상세 불렛
    items_r = [
        ("로직 개선", "LiveView.tsx 및 videoFrames.ts에서 정수 프레임 인덱스 중심으로 이동 및 Seek 처리 통일"),
        ("동기화 보장", "실제 재생 시간을 대상 프레임 구간 중앙값으로 Seek, 이동 중 Grab 차단"),
        ("E2E 재검증", "실제 Chrome TC-01 PASS (120→121→120 정상 복귀 및 0·479 경계값 확인)"),
        ("단위 테스트", "프레임 인덱스 변환·경계 조건 관련 단위 테스트 8건 전원 PASS 완료")
    ]
    tb_r_text = s1.shapes.add_textbox(right_x1 + Inches(0.18), card_y1 + Inches(1.42), col_w1 - Inches(0.36), card_h1 - Inches(1.5))
    tf_rt = tb_r_text.text_frame
    tf_rt.word_wrap = True
    for i, (k, v) in enumerate(items_r):
        p = tf_rt.paragraphs[0] if i == 0 else tf_rt.add_paragraph()
        p.text = "✔ " + k + ": " + v
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(4) if i > 0 else Pt(0)

    # 하단 시사점 배너
    add_card(s1, Inches(0.8), Inches(6.52), Inches(11.733), Inches(0.62), bg_color=CARD_BG_ALT, border_color=PRIMARY_CYAN)
    tb_bot1 = s1.shapes.add_textbox(Inches(0.95), Inches(6.58), Inches(11.4), Inches(0.5))
    tf_b1 = tb_bot1.text_frame
    tf_b1.word_wrap = True
    p_b1 = tf_b1.paragraphs[0]
    p_b1.text = "💡 엔지니어링 의의: 수동 점검에서 놓치기 쉬운 1프레임 미세 오차를 E2E 자동화로 정확히 포착하고, 코드 수정 후 동일 테스트로 즉시 회귀 검증을 완성하여 표시 프레임과 Grab 대상의 일치성을 완벽히 확보했습니다."
    p_b1.font.name = FONT_FAMILY
    p_b1.font.size = Pt(9.5)
    p_b1.font.color.rgb = TEXT_WHITE

    # ==========================================
    # SLIDE 2: 테스트 현황 및 데이터 보호 기반 후속 계획
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "자동화 테스트 현황 및 데이터 보호 기반 검증 계획",
               category_text="VISIONINSIGHT  |  TEST STATUS & DATA INTEGRITY GUARDRAILS",
               subtitle_text="누적 3 PASS / 2 BLOCKED 현황 · Capture 소유권 및 TTL 정리 위험 선제 차단 · 격리 환경 로드맵")

    # 발표자 노트 (약 30초 대본)
    s2.notes_slide.notes_text_frame.text = (
        "현재 읽기 전용 자동화 흐름 5개 중 3개 기능을 검증했고, Edge와 영역별 밝기 2개는 자동 검증 대기 상태입니다. "
        "기능이 고장 난 것이 아니라 기존 Grab 이미지의 세션 소유권 때문에 별도 브라우저에서 접근할 수 없었습니다. "
        "또한 테스트용 Capture 생성 과정에서 오래된 파일이 정리될 위험을 미리 발견해 실행을 차단했습니다. "
        "발표 이후 이번 주에는 격리된 테스트 환경에서 남은 두 흐름을 검증할 계획입니다."
    )

    card_h2_top = Inches(3.45)
    card_y2_top = Inches(1.35)

    # 좌측: 3 PASS / 2 BLOCKED 현황
    add_card(s2, left_x1, card_y2_top, col_w1, card_h2_top, border_color=PRIMARY_CYAN)
    tb_s2_l = s2.shapes.add_textbox(left_x1 + Inches(0.18), card_y2_top + Inches(0.12), col_w1 - Inches(0.36), card_h2_top - Inches(0.2))
    tf_s2_l = tb_s2_l.text_frame
    tf_s2_l.word_wrap = True

    p_s2l = tf_s2_l.paragraphs[0]
    p_s2l.text = "1단계 읽기 전용 검증 현황: 3 PASS / 2 BLOCKED"
    p_s2l.font.name = FONT_FAMILY
    p_s2l.font.size = Pt(13)
    p_s2l.font.bold = True
    p_s2l.font.color.rgb = PRIMARY_CYAN

    p_pass_h = tf_s2_l.add_paragraph()
    p_pass_h.text = "✔ 자동화 검증 완료 (3 PASS)"
    p_pass_h.font.name = FONT_FAMILY
    p_pass_h.font.size = Pt(11)
    p_pass_h.font.bold = True
    p_pass_h.font.color.rgb = SUCCESS_GREEN
    p_pass_h.space_before = Pt(6)

    pass_items = [
        ("TC-01 (PASS)", "비디오 재생·Pause·Seek 및 ±1 Frame 조작"),
        ("TC-02 (PASS)", "검사 상위·하위 탭 전환 시 기존 검사 ID 유지 (불필요 요청 0회)"),
        ("TC-03 (PASS)", "Heatmap·Zoom·Pan·Pixel Inspector (INSP-001651 원본 대조)")
    ]
    for code, desc in pass_items:
        p = tf_s2_l.add_paragraph()
        p.text = "   • " + code + ": " + desc
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(2)

    p_blk_h = tf_s2_l.add_paragraph()
    p_blk_h.text = "⏸ 자동 검증 보류 (2 BLOCKED) — * 기능 오류 아님, 데이터 보호"
    p_blk_h.font.name = FONT_FAMILY
    p_blk_h.font.size = Pt(11)
    p_blk_h.font.bold = True
    p_blk_h.font.color.rgb = WARNING_AMBER
    p_blk_h.space_before = Pt(8)

    blk_items = [
        ("TC-04 (BLOCKED)", "병 외곽 / 내부 림 Edge 오버레이 및 신뢰도 자동 단언"),
        ("TC-05 (BLOCKED)", "A~D 영역별 GV 통계 및 계산 불가 상태 자동 단언")
    ]
    for code, desc in blk_items:
        p = tf_s2_l.add_paragraph()
        p.text = "   • " + code + ": " + desc
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(2)

    p_note = tf_s2_l.add_paragraph()
    p_note.text = "   ※ 원래 사용자 Chrome에서는 Edge 및 영역별 GV 정상 표시 수동 확인 완료"
    p_note.font.name = FONT_FAMILY
    p_note.font.size = Pt(8.5)
    p_note.font.color.rgb = TEXT_MUTED
    p_note.space_before = Pt(4)

    # 우측: 원인 분석 및 선제적 삭제 위험 차단
    add_card(s2, right_x1, card_y2_top, col_w1, card_h2_top, border_color=WARNING_AMBER)
    tb_s2_r = s2.shapes.add_textbox(right_x1 + Inches(0.18), card_y2_top + Inches(0.12), col_w1 - Inches(0.36), card_h2_top - Inches(0.2))
    tf_s2_r = tb_s2_r.text_frame
    tf_s2_r.word_wrap = True

    p_s2r = tf_s2_r.paragraphs[0]
    p_s2r.text = "원인 규명 및 파일 삭제 위험 선제 차단 (안전 가드레일)"
    p_s2r.font.name = FONT_FAMILY
    p_s2r.font.size = Pt(13)
    p_s2r.font.bold = True
    p_s2r.font.color.rgb = WARNING_AMBER

    reasons = [
        ("1. Capture 소유 세션 분리 원칙", "Playwright가 별도 Chrome 컨텍스트로 실행되어 기존 INSP-001694 Grab의 세션 소유권 부재. 소유권 우회나 Grab 출처 위장 없이 정직하게 SKIP 처리"),
        ("2. TTL 정리 위험 사전 발견", "신규 세션 발급 과정에서 bootstrap→cleanup이 동작해 30분 초과된 기존 PNG가 삭제될 수 있는 잠재 위험을 사전 검증 단계에서 포착"),
        ("3. 선제적 안전 조치 (차단)", "무리하게 실행하지 않고 보존 게이트에서 신규 세션·Grab 실행을 즉시 차단 (세션 POST 0회 / Grab 0회, 기존 PNG 및 감사 로그 100% 무결점 보존)"),
        ("4. 안전 검증 완료", "안전 보존 게이트 단위 테스트 5 PASS 완료 (실제 삭제 사고 0건)")
    ]
    for r_title, r_desc in reasons:
        p = tf_s2_r.add_paragraph()
        p.text = "• " + r_title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(5)

        pd = tf_s2_r.add_paragraph()
        pd.text = "   " + r_desc
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_LIGHT
        pd.space_before = Pt(1)

    # 하단 와이드 카드: 이번 주 격리 테스트 환경 실행 계획
    add_card(s2, Inches(0.8), Inches(4.95), Inches(11.733), Inches(2.2), border_color=SUCCESS_GREEN)
    tb_bot2 = s2.shapes.add_textbox(Inches(0.95), Inches(5.02), Inches(11.4), Inches(2.05))
    tf_b2 = tb_bot2.text_frame
    tf_b2.word_wrap = True

    pb2_title = tf_b2.paragraphs[0]
    pb2_title.text = "이번 주 실행 로드맵: 격리 테스트 환경(Isolated Test Bed) 기반 자동화 완결"
    pb2_title.font.name = FONT_FAMILY
    pb2_title.font.size = Pt(13)
    pb2_title.font.bold = True
    pb2_title.font.color.rgb = SUCCESS_GREEN

    # 2단 분기 박스 (좌: 기존 서버 보존 / 우: 격리 테스트 서버)
    # Box 1: 기존 서버 보존
    sub_box_1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(5.35), Inches(5.5), Inches(1.65))
    sub_box_1.fill.solid()
    sub_box_1.fill.fore_color.rgb = CARD_BG_ALT
    sub_box_1.line.color.rgb = PRIMARY_CYAN
    sub_box_1.line.width = Pt(1)

    tf_sb1 = sub_box_1.text_frame
    tf_sb1.word_wrap = True
    p_sb1 = tf_sb1.paragraphs[0]
    p_sb1.text = "[현재 운영] 기존 서버 및 실데이터 절대 보존"
    p_sb1.font.name = FONT_FAMILY
    p_sb1.font.size = Pt(11)
    p_sb1.font.bold = True
    p_sb1.font.color.rgb = PRIMARY_CYAN

    sb1_items = [
        "실서버 inspections.jsonl 및 기존 Capture PNG 무단 수정·삭제 원천 금지",
        "발표 시연 환경 및 검증 데이터의 영구적 안정성 확보",
        "운영 중인 백엔드(포트 8000) 프로세스 무중단 유지"
    ]
    for it in sb1_items:
        p = tf_sb1.add_paragraph()
        p.text = "• " + it
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(3)

    # Box 2: 격리 테스트 서버
    sub_box_2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(5.35), Inches(5.5), Inches(1.65))
    sub_box_2.fill.solid()
    sub_box_2.fill.fore_color.rgb = CARD_BG_ALT
    sub_box_2.line.color.rgb = SUCCESS_GREEN
    sub_box_2.line.width = Pt(1)

    tf_sb2 = sub_box_2.text_frame
    tf_sb2.word_wrap = True
    p_sb2 = tf_sb2.paragraphs[0]
    p_sb2.text = "[이번 주 실행] 격리 테스트 환경 기반 TC-04·05 완주"
    p_sb2.font.name = FONT_FAMILY
    p_sb2.font.size = Pt(11)
    p_sb2.font.bold = True
    p_sb2.font.color.rgb = SUCCESS_GREEN

    sb2_items = [
        "별도 테스트 포트 + 빈 임시 Capture 디렉터리의 독립 테스트 서버 기동",
        "정상 소유 세션으로 1회 Grab 생성 ➔ Edge 및 영역별 GV 자동 검증 완결",
        "Gemini Computer Use 시각 인지 브리지 결합 (레이아웃/오버레이 검증 보강)"
    ]
    for it in sb2_items:
        p = tf_sb2.add_paragraph()
        p.text = "✔ " + it
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(3)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Test automation slides successfully created at: {output_path}")

if __name__ == "__main__":
    create_test_automation_slides()
